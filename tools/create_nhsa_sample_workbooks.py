from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook


SOURCES = [
    (
        Path("D:/20260320发布-医保医用耗材分类与代码数据库更新-14737190条/16、医保医用耗材代码_C09_体外循环材料_全量规格型号信息.xlsx"),
        Path("data/samples/external/nhsa/H-UDMP-SAMPLE-NHSA-C09-FULL-SPEC-v1.0.xlsx"),
        12,
    ),
    (
        Path("D:/20260320发布-医保医用耗材分类与代码数据库更新-14737190条/30、停用表.xlsx"),
        Path("data/samples/external/nhsa/H-UDMP-SAMPLE-NHSA-DISABLED-v1.0.xlsx"),
        12,
    ),
    (
        Path("D:/20260320发布-医保医用耗材分类与代码数据库更新-14737190条/31、转码表.xlsx"),
        Path("data/samples/external/nhsa/H-UDMP-SAMPLE-NHSA-TRANSCODE-v1.0.xlsx"),
        12,
    ),
]


def copy_sample(source: Path, target: Path, max_rows: int) -> None:
    source_wb = load_workbook(source, read_only=True, data_only=True)
    source_ws = source_wb["Query1"] if "Query1" in source_wb.sheetnames else source_wb[source_wb.sheetnames[0]]

    target_wb = Workbook()
    target_ws = target_wb.active
    target_ws.title = "Query1"

    for row_index, row in enumerate(source_ws.iter_rows(values_only=True), start=1):
        target_ws.append(list(row))
        if row_index >= max_rows + 1:
            break

    source_wb.close()
    target.parent.mkdir(parents=True, exist_ok=True)
    target_wb.save(target)


def main() -> None:
    repo_root = Path(__file__).resolve().parents[1]
    for source, relative_target, row_count in SOURCES:
        target = repo_root / relative_target
        if not source.exists():
            raise FileNotFoundError(source)
        copy_sample(source, target, row_count)
        print(f"created {target}")


if __name__ == "__main__":
    main()
