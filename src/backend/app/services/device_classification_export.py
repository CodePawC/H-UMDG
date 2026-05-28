"""Device classification catalog export task service."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import uuid
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.models.tables import (
    CatalogCorrectionOrder,
    DeviceClassificationCatalogVersion,
    DeviceClassificationImportAuditLog,
    DeviceClassificationSourceFile,
    RefDeviceClassificationCatalog,
    RefDeviceClassificationRevision,
    SysImportBatch,
)


EXPORT_SCOPE_LABELS = {
    "effective": "当前有效目录",
    "version": "指定版本目录",
    "batch": "指定导入批次",
    "category": "指定分类范围",
    "abnormal": "异常待确认数据",
    "historical": "已作废已合并数据",
    "version_diff": "版本差异数据",
}

EXPORT_TEMPLATE_LABELS = {
    "standard": "标准目录模板",
    "compact": "简洁目录模板",
    "governance": "数据治理模板",
    "version_diff": "版本差异模板",
    "interface": "系统接口模板",
    "regulatory": "监管备查模板",
}

EXPORT_FORMAT_LABELS = {
    "xlsx": "Excel .xlsx",
    "csv": "CSV .csv",
    "json": "JSON .json",
    "zip": "ZIP 压缩包",
}

FIELD_DEFINITIONS = {
    "catalog_code": ("分类编码", "国家医疗器械分类目录编码。"),
    "major_category": ("一级分类", "一级分类编码与名称。"),
    "level_1_category": ("二级分类", "二级分类编码与名称。"),
    "catalog_name": ("目录条目", "目录条目名称。"),
    "product_description": ("产品描述", "产品描述字段。"),
    "intended_use": ("预期用途", "预期用途字段。"),
    "product_examples": ("品名举例", "品名举例字段。"),
    "management_class": ("管理类别", "医疗器械管理类别。"),
    "level": ("目录层级", "按编码和分类字段推断的目录层级。"),
    "data_status": ("数据状态", "主数据治理状态。"),
    "source_batch_id": ("来源批次", "数据来源导入批次号。"),
    "source_file_name": ("来源文件", "原始来源文件名。"),
    "source_position": ("原始位置", "原始文件行号或位置。"),
    "validation_result": ("校验结果", "导入核验和解析结果。"),
    "correction_status": ("纠错状态", "关联纠错单最新状态。"),
    "change_type": ("变更类型", "新增、修订、废止等变更类型。"),
    "before_data": ("变更前内容", "版本差异中的变更前字段快照。"),
    "after_data": ("变更后内容", "版本差异中的变更后字段快照。"),
    "revision_reason": ("修订依据", "版本修订依据或说明。"),
    "effective_status": ("生效状态", "差异记录生效状态。"),
    "catalog_id": ("主键ID", "系统内部目录主键。"),
    "parent_id": ("父级ID", "系统内部父级目录主键。"),
    "version_id": ("版本ID", "目录版本标识。"),
    "status": ("状态", "系统状态字段。"),
    "updated_at": ("更新时间", "目录记录最近更新时间。"),
    "source_announcement": ("来源公告", "来源公告或标准来源。"),
    "source_file_hash": ("文件 Hash", "来源文件 SHA256。"),
    "imported_at": ("导入时间", "来源批次导入时间。"),
    "operator_name": ("操作人", "来源批次操作人。"),
    "audit_status": ("审计状态", "来源和导出审计状态。"),
}

TEMPLATE_FIELDS = {
    "standard": [
        "catalog_code",
        "major_category",
        "level_1_category",
        "catalog_name",
        "product_description",
        "intended_use",
        "product_examples",
        "management_class",
    ],
    "compact": ["catalog_code", "level", "catalog_name", "management_class"],
    "governance": [
        "catalog_code",
        "major_category",
        "level_1_category",
        "catalog_name",
        "product_description",
        "intended_use",
        "product_examples",
        "management_class",
        "data_status",
        "source_batch_id",
        "source_file_name",
        "source_position",
        "validation_result",
        "correction_status",
    ],
    "version_diff": [
        "change_type",
        "catalog_code",
        "before_data",
        "after_data",
        "revision_reason",
        "source_batch_id",
        "effective_status",
    ],
    "interface": ["catalog_id", "parent_id", "catalog_code", "catalog_name", "level", "version_id", "status", "updated_at"],
    "regulatory": [
        "catalog_code",
        "major_category",
        "level_1_category",
        "catalog_name",
        "product_description",
        "intended_use",
        "product_examples",
        "management_class",
        "source_announcement",
        "source_file_name",
        "source_file_hash",
        "imported_at",
        "operator_name",
        "audit_status",
    ],
}


@dataclass
class ExportBuildResult:
    record: dict[str, Any]
    path: Path


def export_root() -> Path:
    root = Path.cwd() / "runtime" / "device-classification-exports"
    root.mkdir(parents=True, exist_ok=True)
    return root


def _records_path() -> Path:
    return export_root() / "records.json"


def _read_records() -> list[dict[str, Any]]:
    path = _records_path()
    if not path.exists():
        return []
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []
    return payload if isinstance(payload, list) else []


def _write_records(records: list[dict[str, Any]]) -> None:
    _records_path().write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


def list_export_records() -> list[dict[str, Any]]:
    return sorted(_read_records(), key=lambda item: str(item.get("created_at") or ""), reverse=True)


def get_export_record(task_no: str) -> dict[str, Any] | None:
    return next((item for item in _read_records() if item.get("task_no") == task_no), None)


def update_export_record(record: dict[str, Any]) -> None:
    records = _read_records()
    next_records = [item for item in records if item.get("task_no") != record.get("task_no")]
    next_records.append(record)
    _write_records(next_records)


def export_templates() -> list[dict[str, Any]]:
    return [
        {
            "template_id": template_id,
            "template_name": label,
            "fields": [
                {"field": field, "label": FIELD_DEFINITIONS[field][0], "description": FIELD_DEFINITIONS[field][1]}
                for field in TEMPLATE_FIELDS[template_id]
            ],
            "sensitive_fields": [
                field for field in TEMPLATE_FIELDS[template_id] if field in {"catalog_id", "parent_id", "source_file_hash", "audit_status"}
            ],
        }
        for template_id, label in EXPORT_TEMPLATE_LABELS.items()
    ]


def _next_task_no() -> str:
    date_part = datetime.now(UTC).strftime("%Y%m%d")
    today_count = sum(1 for item in _read_records() if str(item.get("task_no") or "").startswith(f"EXP{date_part}"))
    return f"EXP{date_part}{today_count + 1:04d}"


def _safe_filename(value: str) -> str:
    normalized = re.sub(r'[\\/:*?"<>|\s]+', "_", value).strip("_")
    return normalized[:180] or "device-classification-export"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _iso(value: Any) -> str | None:
    return value.isoformat() if value else None


def _catalog_code(row: RefDeviceClassificationCatalog) -> str:
    return row.category_no or row.level_2_category_no or row.level_1_category_no or row.major_category_no or ""


def _catalog_name(row: RefDeviceClassificationCatalog) -> str:
    return row.level_2_category or row.level_1_category or row.major_category_name or ""


def _level(row: RefDeviceClassificationCatalog) -> str:
    if row.level_2_category or row.level_2_category_no:
        return "三级"
    if row.level_1_category or row.level_1_category_no:
        return "二级"
    return "一级"


def _source_position(row: RefDeviceClassificationCatalog) -> str:
    if row.row_number:
        return f"第 {row.row_number} 行"
    return f"表 {row.source_table_index}"


def _json_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _latest_version_label(db: Session) -> str:
    version = db.scalar(
        select(DeviceClassificationCatalogVersion)
        .where(DeviceClassificationCatalogVersion.is_current.is_(True))
        .order_by(DeviceClassificationCatalogVersion.created_at.desc())
        .limit(1)
    )
    if version:
        return version.version_no
    return "NMPA-2017-104"


def _load_context(db: Session, batch_ids: set[str], catalog_ids: set[uuid.UUID]) -> dict[str, Any]:
    source_files = {
        item.batch_id: item
        for item in db.scalars(select(DeviceClassificationSourceFile).where(DeviceClassificationSourceFile.batch_id.in_(batch_ids))).all()
    } if batch_ids else {}
    batches = {
        item.batch_id: item
        for item in db.scalars(select(SysImportBatch).where(SysImportBatch.batch_id.in_(batch_ids))).all()
    } if batch_ids else {}
    corrections: dict[str, str] = {}
    if catalog_ids:
        rows = db.execute(
            select(CatalogCorrectionOrder.catalog_id, CatalogCorrectionOrder.status)
            .where(CatalogCorrectionOrder.catalog_id.in_(catalog_ids))
            .order_by(CatalogCorrectionOrder.created_at.desc())
        ).all()
        for catalog_id, status in rows:
            corrections.setdefault(str(catalog_id), status)
    return {"source_files": source_files, "batches": batches, "corrections": corrections}


def _catalog_query(db: Session, criteria: dict[str, Any]):
    scope_type = str(criteria.get("scope_type") or "effective")
    query = select(RefDeviceClassificationCatalog)
    if scope_type == "effective":
        query = query.where(
            RefDeviceClassificationCatalog.data_status == "effective",
            RefDeviceClassificationCatalog.hidden_in_tree.is_(False),
        )
    elif scope_type == "abnormal":
        query = query.where(RefDeviceClassificationCatalog.data_status.in_(["pending_confirm", "parse_abnormal", "corrected"]))
    elif scope_type == "historical":
        query = query.where(RefDeviceClassificationCatalog.data_status.in_(["deprecated", "merged", "rollbacked"]))
    elif scope_type == "batch" and criteria.get("batch_id"):
        query = query.where(RefDeviceClassificationCatalog.batch_id == str(criteria["batch_id"]))
    elif scope_type == "version" and criteria.get("version_id"):
        try:
            version_uuid = uuid.UUID(str(criteria["version_id"]))
        except ValueError:
            version_uuid = None
        if version_uuid:
            version = db.get(DeviceClassificationCatalogVersion, version_uuid)
            if version:
                query = query.where(RefDeviceClassificationCatalog.batch_id == version.batch_id)
    elif scope_type == "category" and criteria.get("category_keyword"):
        keyword = f"%{str(criteria['category_keyword']).strip()}%"
        query = query.where(
            or_(
                RefDeviceClassificationCatalog.major_category_name.ilike(keyword),
                RefDeviceClassificationCatalog.level_1_category.ilike(keyword),
                RefDeviceClassificationCatalog.level_2_category.ilike(keyword),
                RefDeviceClassificationCatalog.major_category_no.ilike(keyword),
                RefDeviceClassificationCatalog.level_1_category_no.ilike(keyword),
                RefDeviceClassificationCatalog.level_2_category_no.ilike(keyword),
                RefDeviceClassificationCatalog.category_no.ilike(keyword),
            )
        )
    return query.order_by(
        RefDeviceClassificationCatalog.major_category_no.nulls_last(),
        RefDeviceClassificationCatalog.level_1_category_no.nulls_last(),
        RefDeviceClassificationCatalog.level_2_category_no.nulls_last(),
        RefDeviceClassificationCatalog.category_no.nulls_last(),
        RefDeviceClassificationCatalog.row_number.asc(),
    )


def _revision_rows(db: Session, criteria: dict[str, Any]) -> list[dict[str, Any]]:
    query = select(RefDeviceClassificationRevision)
    if criteria.get("batch_id"):
        query = query.where(RefDeviceClassificationRevision.batch_id == str(criteria["batch_id"]))
    rows = db.scalars(query.order_by(RefDeviceClassificationRevision.created_at.desc()).limit(5000)).all()
    result: list[dict[str, Any]] = []
    for row in rows:
        after_payload = row.new_payload or {}
        before_payload = row.old_payload or {}
        result.append(
            {
                "change_type": row.change_type,
                "catalog_code": after_payload.get("category_no") or before_payload.get("category_no") or "",
                "before_data": before_payload,
                "after_data": after_payload,
                "revision_reason": row.reason,
                "source_batch_id": row.batch_id,
                "effective_status": "已记录",
            }
        )
    return result


def _catalog_rows(db: Session, criteria: dict[str, Any]) -> list[dict[str, Any]]:
    rows = db.scalars(_catalog_query(db, criteria).limit(10000)).all()
    context = _load_context(db, {row.batch_id for row in rows}, {row.catalog_id for row in rows})
    source_files = context["source_files"]
    batches = context["batches"]
    corrections = context["corrections"]
    current_version = _latest_version_label(db)
    result: list[dict[str, Any]] = []
    for row in rows:
        source_file = source_files.get(row.batch_id)
        batch = batches.get(row.batch_id)
        result.append(
            {
                "catalog_id": str(row.catalog_id),
                "parent_id": "",
                "catalog_code": _catalog_code(row),
                "major_category": " ".join(part for part in [row.major_category_no, row.major_category_name] if part),
                "level_1_category": " ".join(part for part in [row.level_1_category_no, row.level_1_category] if part),
                "catalog_name": _catalog_name(row),
                "product_description": row.product_description,
                "intended_use": row.intended_use,
                "product_examples": row.product_examples,
                "management_class": row.management_class,
                "level": _level(row),
                "data_status": row.data_status,
                "source_batch_id": row.batch_id,
                "source_file_name": row.source_file_name,
                "source_position": _source_position(row),
                "validation_result": row.status_reason or "通过",
                "correction_status": corrections.get(str(row.catalog_id), "无纠错单"),
                "version_id": current_version,
                "status": row.data_status,
                "updated_at": _iso(row.updated_at),
                "source_announcement": (batch.catalog_version if batch else None) or "国家药监局医疗器械分类目录",
                "source_file_hash": source_file.sha256 if source_file else None,
                "imported_at": _iso(batch.finished_at or batch.created_at) if batch else _iso(row.created_at),
                "operator_name": batch.operator_name if batch else None,
                "audit_status": "已留痕",
            }
        )
    return result


def _resolve_fields(template_id: str, requested_fields: list[str] | None) -> list[str]:
    default_fields = TEMPLATE_FIELDS.get(template_id, TEMPLATE_FIELDS["standard"])
    if not requested_fields:
        return list(default_fields)
    allowed = set(default_fields)
    return [field for field in requested_fields if field in allowed and field in FIELD_DEFINITIONS] or list(default_fields)


def _rows_for_export(db: Session, criteria: dict[str, Any], template_id: str) -> list[dict[str, Any]]:
    if str(criteria.get("scope_type") or "") == "version_diff" or template_id == "version_diff":
        return _revision_rows(db, criteria)
    return _catalog_rows(db, criteria)


def estimate_export_count(db: Session, criteria: dict[str, Any], template_id: str = "standard") -> int:
    if str(criteria.get("scope_type") or "") == "version_diff" or template_id == "version_diff":
        query = select(func.count()).select_from(RefDeviceClassificationRevision)
        if criteria.get("batch_id"):
            query = query.where(RefDeviceClassificationRevision.batch_id == str(criteria["batch_id"]))
        return int(db.scalar(query) or 0)
    count_query = select(func.count()).select_from(_catalog_query(db, criteria).subquery())
    return int(db.scalar(count_query) or 0)


def _summary_payload(
    *,
    task_no: str,
    criteria: dict[str, Any],
    template_id: str,
    format_id: str,
    row_count: int,
    fields: list[str],
    operator_name: str,
    created_at: datetime,
    purpose: str | None,
) -> dict[str, Any]:
    return {
        "数据域": "医疗器械分类目录",
        "导出任务号": task_no,
        "导出范围": EXPORT_SCOPE_LABELS.get(str(criteria.get("scope_type") or "effective"), str(criteria.get("scope_type") or "")),
        "目录版本": str(criteria.get("version_label") or criteria.get("version_id") or _json_cell(criteria.get("batch_id")) or "当前启用版本"),
        "导出模板": EXPORT_TEMPLATE_LABELS.get(template_id, template_id),
        "导出格式": EXPORT_FORMAT_LABELS.get(format_id, format_id),
        "导出时间": created_at.isoformat(),
        "导出人": operator_name,
        "数据来源": "H-UMDG 主数据治理库",
        "导出条目数": row_count,
        "导出用途": purpose or "目录复核",
        "注意事项": "导出文件仅代表生成时点的数据快照，敏感字段按当前角色权限控制。",
        "导出字段": "、".join(FIELD_DEFINITIONS[field][0] for field in fields),
    }


def _write_xlsx(
    path: Path,
    rows: list[dict[str, Any]],
    fields: list[str],
    summary: dict[str, Any],
    *,
    include_governance_sheet: bool,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "导出说明"
    sheet.append(["项目", "内容"])
    for key, value in summary.items():
        sheet.append([key, _json_cell(value)])

    data_sheet = workbook.create_sheet("目录数据")
    data_sheet.append([FIELD_DEFINITIONS[field][0] for field in fields])
    for row in rows:
        data_sheet.append([_json_cell(row.get(field)) for field in fields])

    fields_sheet = workbook.create_sheet("字段说明")
    fields_sheet.append(["字段编码", "字段名称", "说明"])
    for field in fields:
        label, description = FIELD_DEFINITIONS[field]
        fields_sheet.append([field, label, description])

    source_sheet = workbook.create_sheet("来源与版本")
    source_sheet.append(["来源批次", "来源文件", "文件 Hash", "导入时间", "操作人", "版本"])
    seen: set[tuple[str, str]] = set()
    for row in rows:
        key = (str(row.get("source_batch_id") or ""), str(row.get("source_file_name") or ""))
        if key in seen:
            continue
        seen.add(key)
        source_sheet.append([
            row.get("source_batch_id"),
            row.get("source_file_name"),
            row.get("source_file_hash"),
            row.get("imported_at"),
            row.get("operator_name"),
            row.get("version_id"),
        ])

    if include_governance_sheet:
        governance_sheet = workbook.create_sheet("异常与纠错记录")
        governance_sheet.append(["分类编码", "目录条目", "数据状态", "校验结果", "纠错状态", "来源批次", "原始位置"])
        for row in rows:
            if row.get("data_status") != "effective" or row.get("correction_status") not in {None, "", "无纠错单"}:
                governance_sheet.append([
                    row.get("catalog_code"),
                    row.get("catalog_name"),
                    row.get("data_status"),
                    row.get("validation_result"),
                    row.get("correction_status"),
                    row.get("source_batch_id"),
                    row.get("source_position"),
                ])

    for worksheet in workbook.worksheets:
        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 50)
    workbook.save(path)


def _write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow([FIELD_DEFINITIONS[field][0] for field in fields])
        for row in rows:
            writer.writerow([_json_cell(row.get(field)) for field in fields])


def _write_json(path: Path, rows: list[dict[str, Any]], fields: list[str], summary: dict[str, Any]) -> None:
    payload = {
        "summary": summary,
        "fields": [{"field": field, "label": FIELD_DEFINITIONS[field][0]} for field in fields],
        "items": [{field: row.get(field) for field in fields} for row in rows],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_zip(path: Path, rows: list[dict[str, Any]], fields: list[str], summary: dict[str, Any], *, include_governance_sheet: bool) -> None:
    temp_xlsx = path.with_suffix(".xlsx")
    temp_json = path.with_suffix(".json")
    temp_readme = path.with_name(f"{path.stem}-导出说明.txt")
    try:
        _write_xlsx(temp_xlsx, rows, fields, summary, include_governance_sheet=include_governance_sheet)
        _write_json(temp_json, rows, fields, summary)
        temp_readme.write_text(
            "\n".join(f"{key}: {_json_cell(value)}" for key, value in summary.items()),
            encoding="utf-8",
        )
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.write(temp_xlsx, arcname=temp_xlsx.name)
            archive.write(temp_json, arcname=temp_json.name)
            archive.write(temp_readme, arcname="导出说明.txt")
    finally:
        for temp_path in (temp_xlsx, temp_json, temp_readme):
            if temp_path.exists():
                temp_path.unlink()


def create_export_task(
    db: Session,
    payload: dict[str, Any],
    *,
    operator_name: str,
    operator_role: str | None,
    client_ip: str | None,
    user_agent: str | None,
) -> ExportBuildResult:
    format_id = str(payload.get("format") or "xlsx").lower()
    if format_id not in EXPORT_FORMAT_LABELS:
        raise ValueError("EXPORT_FORMAT_NOT_SUPPORTED")
    template_id = str(payload.get("template_id") or "standard")
    if template_id not in TEMPLATE_FIELDS:
        raise ValueError("EXPORT_TEMPLATE_NOT_FOUND")

    criteria = dict(payload.get("criteria") or {})
    criteria.setdefault("scope_type", payload.get("scope_type") or "effective")
    fields = _resolve_fields(template_id, payload.get("fields") if isinstance(payload.get("fields"), list) else None)
    rows = _rows_for_export(db, criteria, template_id)
    task_no = _next_task_no()
    created_at = datetime.now(UTC)
    purpose = str(payload.get("purpose") or "").strip() or None
    task_name = str(payload.get("task_name") or "").strip() or f"{EXPORT_SCOPE_LABELS.get(criteria['scope_type'], '目录')}{EXPORT_TEMPLATE_LABELS.get(template_id, '')}导出"
    version_label = str(payload.get("version_label") or criteria.get("version_label") or _latest_version_label(db))
    criteria["version_label"] = version_label
    summary = _summary_payload(
        task_no=task_no,
        criteria=criteria,
        template_id=template_id,
        format_id=format_id,
        row_count=len(rows),
        fields=fields,
        operator_name=operator_name,
        created_at=created_at,
        purpose=purpose,
    )
    file_stem = _safe_filename(f"医疗器械分类目录_{summary['导出范围']}_{version_label}_{created_at:%Y%m%d}_{task_no}")
    file_name = f"{file_stem}.{format_id}"
    path = export_root() / file_name

    if format_id == "xlsx":
        _write_xlsx(path, rows, fields, summary, include_governance_sheet=template_id in {"governance", "regulatory"})
    elif format_id == "csv":
        _write_csv(path, rows, fields)
    elif format_id == "zip":
        _write_zip(path, rows, fields, summary, include_governance_sheet=template_id in {"governance", "regulatory"})
    else:
        _write_json(path, rows, fields, summary)

    file_hash = _sha256(path)
    audit_item = {
        "action": "EXPORT_CREATED",
        "operator_name": operator_name,
        "operator_role": operator_role,
        "occurred_at": created_at.isoformat(),
        "client_ip": client_ip,
        "user_agent": user_agent,
        "export_scope": summary["导出范围"],
        "export_fields": fields,
        "export_format": format_id,
        "export_purpose": purpose,
        "file_hash": file_hash,
    }
    record = {
        "task_no": task_no,
        "task_name": task_name,
        "scope_type": criteria.get("scope_type"),
        "scope_label": summary["导出范围"],
        "version_label": version_label,
        "template_id": template_id,
        "template_name": EXPORT_TEMPLATE_LABELS.get(template_id, template_id),
        "format": format_id,
        "format_label": EXPORT_FORMAT_LABELS.get(format_id, format_id),
        "row_count": len(rows),
        "operator_name": operator_name,
        "operator_role": operator_role,
        "created_at": created_at.isoformat(),
        "purpose": purpose,
        "file_status": "已生成",
        "file_name": file_name,
        "file_path": str(path),
        "file_size_bytes": path.stat().st_size,
        "file_hash": file_hash,
        "download_count": 0,
        "last_downloaded_at": None,
        "criteria": criteria,
        "fields": fields,
        "summary": summary,
        "download_logs": [],
        "audit_logs": [audit_item],
    }
    records = _read_records()
    records.append(record)
    _write_records(records)
    db.add(
        DeviceClassificationImportAuditLog(
            batch_id=task_no,
            action="CATALOG_EXPORT_CREATED",
            operator_name=operator_name,
            message=f"生成医疗器械分类目录导出任务 {task_no}",
            payload=audit_item,
        )
    )
    return ExportBuildResult(record=record, path=path)


def record_export_download(
    db: Session,
    record: dict[str, Any],
    *,
    operator_name: str,
    client_ip: str | None,
    user_agent: str | None,
) -> dict[str, Any]:
    now = datetime.now(UTC)
    log = {
        "action": "EXPORT_DOWNLOADED",
        "operator_name": operator_name,
        "occurred_at": now.isoformat(),
        "client_ip": client_ip,
        "user_agent": user_agent,
        "file_hash": record.get("file_hash"),
        "result_status": "success",
    }
    record["download_count"] = int(record.get("download_count") or 0) + 1
    record["last_downloaded_at"] = now.isoformat()
    record.setdefault("download_logs", []).append(log)
    record.setdefault("audit_logs", []).append(log)
    update_export_record(record)
    db.add(
        DeviceClassificationImportAuditLog(
            batch_id=str(record.get("task_no") or ""),
            action="CATALOG_EXPORT_DOWNLOADED",
            operator_name=operator_name,
            message=f"下载医疗器械分类目录导出文件 {record.get('task_no')}",
            payload=log,
        )
    )
    return record
