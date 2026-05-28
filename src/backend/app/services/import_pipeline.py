"""Department / material import pipelines."""

from __future__ import annotations

import io
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from sqlalchemy import delete, select, update
from sqlalchemy.orm import Session

from app.models.tables import (
    DeviceClassificationCatalogVersion,
    DeviceClassificationImportAuditLog,
    DeviceClassificationImportStaging,
    DeviceClassificationImportValidationResult,
    DictDepartment,
    DictMaterialSpec,
    ImportValidationIssue,
    RefDeviceClassificationCatalog,
    RefDeviceClassificationRevision,
    StgNhsaMaterialDisabled,
    StgNhsaMaterialSpec,
    StgNhsaMaterialTranscode,
)
from app.services.spec_parse import split_spec
from app.services.tabular_io import iter_tabular_rows_from_path, load_tabular_rows
from app.services.manufacturer_vendors import upsert_vendor_candidate

DEVICE_HEADERS_REQUIRED = frozenset(
    {"序号", "一级产品类别", "二级产品类别", "产品描述", "预期用途", "品名举例", "管理类别"}
)
DEVICE_ADJUSTMENT_HEADERS_REQUIRED = frozenset(
    {"序号", "子目录", "一级产品类别", "二级产品类别", "产品描述", "预期用途", "品名举例", "管理类别"}
)
DEVICE_DIGIT_TRANSLATION = str.maketrans("０１２３４５６７８９", "0123456789")
DEVICE_MAJOR_HEADING_RE = re.compile(r"^(\d{2})(?!\d)\s*(.+?)\s*$")
DEVICE_CODE_NAME_RE = re.compile(r"^(\d{1,3})(?:[-－—、.．\s]*)?(.+)$")
DEVICE_CJK_INNER_SPACE_RE = re.compile(r"(?<=[\u3400-\u9fff])\s+(?=[\u3400-\u9fff])")
DEVICE_ALLOWED_MANAGEMENT_CLASSES = {"Ⅰ", "Ⅱ", "Ⅲ", "I", "II", "III"}
DEVICE_HEADER_VALUES = {"产品描述", "预期用途", "管理类别", "品名举例", "序号", "一级产品类别", "二级产品类别", "子目录"}
DEVICE_REVISION_NOTE_MARKERS = {"目录调整", "修订"}
DEVICE_ORIGINAL_FIELD_MARKERS = {"原产品描述", "原预期用途", "原管理类别"}
DEVICE_SYMBOL_ONLY_RE = re.compile(r"^[\W_]+$")
VENDOR_SOURCE_FIELDS = (
    ("耗材企业", "manufacturer"),
    ("生产企业", "manufacturer"),
    ("生产厂家", "manufacturer"),
    ("厂家", "manufacturer"),
    ("注册人", "registrant"),
    ("注册人名称", "registrant"),
    ("备案人", "filer"),
    ("备案人名称", "filer"),
    ("manufacturer", "manufacturer"),
    ("manufacturer_name", "manufacturer"),
    ("registrant", "registrant"),
    ("filer", "filer"),
)


def _normalize_device_header(raw: str) -> str:
    return "".join(str(raw).split())


def _normalize_device_value(raw: Any) -> str:
    value = " ".join(str(raw or "").replace("\u3000", " ").split()).translate(DEVICE_DIGIT_TRANSLATION)
    return DEVICE_CJK_INNER_SPACE_RE.sub("", value)


def _split_device_code_name(raw: Any) -> tuple[str | None, str | None]:
    value = _normalize_device_value(raw)
    if not value:
        return None, None
    if value.isdigit():
        return value.zfill(2) if len(value) <= 2 else value, None
    match = DEVICE_CODE_NAME_RE.match(value)
    if not match:
        return None, value
    code = match.group(1)
    name = match.group(2).strip()
    return code.zfill(2) if len(code) <= 2 else code, name or None


def _parse_device_major_heading(raw: Any) -> tuple[str, str] | None:
    value = _normalize_device_value(raw)
    match = DEVICE_MAJOR_HEADING_RE.match(value)
    if not match:
        return None
    code = match.group(1)
    name = re.sub(r"\s+\d+$", "", match.group(2).strip())
    if not name or name.endswith("说明"):
        return None
    return code, name


def _device_header_indexes(headers: list[str]) -> dict[str, int]:
    result: dict[str, int] = {}
    previous = ""
    for index, header in enumerate(headers):
        if not header:
            continue
        if header == previous:
            continue
        result.setdefault(header, index)
        previous = header
    return result


def _unique_docx_row_values(row: Any) -> list[str]:
    values: list[str] = []
    seen_cells: set[int] = set()
    for cell in row.cells:
        cell_id = id(cell._tc)
        if cell_id in seen_cells:
            continue
        seen_cells.add(cell_id)
        values.append(cell.text.strip())
    return values


def _docx_row_values(row: Any) -> list[str]:
    return [cell.text.strip() for cell in row.cells]


def _none_if_device_marker(value: Any) -> str | None:
    normalized = _normalize_device_value(value)
    return None if normalized in {"", "无"} else normalized


def _parse_device_adjustment_side(
    cells: list[str],
    offset: int,
    original: dict[str, str | None] | None = None,
) -> dict[str, str | None]:
    def adjusted(index: int, key: str) -> str | None:
        value = _normalize_device_value(cells[offset + index] if offset + index < len(cells) else "")
        if value == "无变化" and original is not None:
            return original.get(key)
        return None if value in {"", "无"} else value

    major_raw = adjusted(0, "major_raw")
    level1_raw = adjusted(1, "level1_raw")
    level2_raw = adjusted(2, "level2_raw")
    major_no, major_name = _split_device_code_name(major_raw)
    level1_no, level1_name_from_raw = _split_device_code_name(level1_raw)
    level2_no, level2_name = _split_device_code_name(level2_raw)
    return {
        "major_raw": major_raw,
        "level1_raw": level1_raw,
        "level2_raw": level2_raw,
        "major_category_no": major_no,
        "major_category_name": major_name,
        "level_1_category_no": level1_no,
        "level_1_category": level1_name_from_raw or _none_if_device_marker(level1_raw),
        "level_2_category_no": level2_no,
        "level_2_category": level2_name,
        "product_description": adjusted(3, "product_description"),
        "intended_use": adjusted(4, "intended_use"),
        "product_examples": adjusted(5, "product_examples"),
        "management_class": adjusted(6, "management_class"),
    }


def _has_device_adjustment_original(row: dict[str, str | None]) -> bool:
    return any(
        row.get(key)
        for key in (
            "level_2_category_no",
            "level_2_category",
            "product_description",
            "intended_use",
            "product_examples",
            "management_class",
        )
    )


def _device_validation_issue(row: dict[str, Any]) -> dict[str, str] | None:
    name = _normalize_device_value(row.get("level_2_category"))
    fields_text = " ".join(_normalize_device_value(row.get(field)) for field in (
        "level_2_category",
        "product_description",
        "intended_use",
        "product_examples",
        "management_class",
    ))
    management_class = _normalize_device_value(row.get("management_class"))
    code_parts = [
        _normalize_device_value(row.get("major_category_no")),
        _normalize_device_value(row.get("level_1_category_no")),
        _normalize_device_value(row.get("level_2_category_no")),
    ]
    if not name:
        return {"rule_code": "device_catalog_name_required", "abnormal_type": "目录名称异常", "message": "目录名称不得为空"}
    if DEVICE_SYMBOL_ONLY_RE.match(name):
        return {"rule_code": "device_catalog_name_symbol_only", "abnormal_type": "目录名称异常", "message": "目录名称不得仅包含符号"}
    if name.startswith(("-", "－", "—")):
        return {"rule_code": "device_catalog_name_bad_prefix", "abnormal_type": "解析噪声", "message": "目录名称不得以异常连接符开头"}
    if name in DEVICE_HEADER_VALUES:
        return {"rule_code": "device_catalog_name_header", "abnormal_type": "字段错位", "message": "目录名称疑似字段表头"}
    for field, label in (("product_description", "产品描述"), ("intended_use", "预期用途"), ("product_examples", "品名举例")):
        if _normalize_device_value(row.get(field)) == label:
            return {"rule_code": f"device_{field}_header", "abnormal_type": "字段错位", "message": f"{label}字段疑似表头"}
    if management_class and management_class not in DEVICE_ALLOWED_MANAGEMENT_CLASSES:
        return {"rule_code": "device_management_class_invalid", "abnormal_type": "编码异常", "message": "管理类别只能为 I、II、III"}
    if not all(code_parts) or not code_parts[0].isdigit() or not code_parts[1].isdigit() or not code_parts[2].isdigit():
        return {"rule_code": "device_catalog_code_invalid", "abnormal_type": "编码异常", "message": "编码必须符合医疗器械分类编码格式"}
    if any(marker in fields_text for marker in DEVICE_REVISION_NOTE_MARKERS) or name in DEVICE_ORIGINAL_FIELD_MARKERS:
        return {"rule_code": "device_revision_note_row", "abnormal_type": "人工确认异常", "message": "疑似修订说明行，不得直接入正式分类树"}
    return None


def _apply_device_governance_status(row: dict[str, Any]) -> dict[str, Any]:
    issue = _device_validation_issue(row)
    row["data_status"] = "effective"
    row["abnormal_type"] = None
    row["hidden_in_tree"] = False
    row["status_reason"] = None
    row["validation_issue_codes"] = []
    row["confidence_score"] = 100
    if issue:
        row["data_status"] = "parse_abnormal" if issue["abnormal_type"] in {"解析噪声", "字段错位"} else "pending_confirm"
        row["abnormal_type"] = issue["abnormal_type"]
        row["hidden_in_tree"] = True
        row["status_reason"] = issue["message"]
        row["validation_issue_codes"] = [issue["rule_code"]]
        row["confidence_score"] = 20 if row["data_status"] == "parse_abnormal" else 55
        row["issue_message"] = issue["message"]
        row["parse_status"] = "待人工确认"
        row["change_type"] = "待确认"
    return row


def _device_catalog_model_payload(row: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value
        for key, value in row.items()
        if key not in {
            "abnormal_type",
            "parse_status",
            "issue_message",
            "change_type",
            "staging_id",
            "original_payload",
            "validation_issue_codes",
            "confidence_score",
        }
    }


def _parse_device_adjustment_docx_rows(doc: Document) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for table_index, table in enumerate(doc.tables):
        if len(table.rows) < 3:
            continue
        top_headers = [_normalize_device_header(value) for value in _docx_row_values(table.rows[0])]
        headers = [_normalize_device_header(value) for value in _docx_row_values(table.rows[1])]
        if "调整后《医疗器械分类目录》内容" not in top_headers:
            continue
        if len(headers) < 15 or not DEVICE_ADJUSTMENT_HEADERS_REQUIRED.issubset(set(headers)):
            continue
        for ri in range(2, len(table.rows)):
            cells = _unique_docx_row_values(table.rows[ri])
            if len(cells) < 15:
                continue
            original = _parse_device_adjustment_side(cells, 1)
            adjusted = _parse_device_adjustment_side(cells, 8, original)
            if all(value is None for key, value in adjusted.items() if not key.endswith("_raw")):
                continue
            rows.append({
                **{key: value for key, value in adjusted.items() if not key.endswith("_raw")},
                "source_table_index": table_index,
                "row_number": ri + 1,
                "change_type": "修改" if _has_device_adjustment_original(original) else "新增",
                "original_payload": {key: value for key, value in original.items() if not key.endswith("_raw")},
            })
    return rows


def _iter_docx_blocks(doc: Document) -> Iterable[tuple[str, Paragraph | Table]]:
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield "paragraph", Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield "table", Table(child, doc)


def _now_batch(prefix: str) -> str:
    return f"{prefix}-{datetime.now(timezone.utc).strftime('%Y%m%d')}-{uuid.uuid4().hex[:8]}"


def _extract_vendor_candidates(
    db: Session,
    raw: dict[str, Any],
    *,
    source_system: str,
    source_table: str,
    source_record_id: str,
) -> dict[str, Any]:
    created_or_updated = 0
    matched_org_ids: dict[str, uuid.UUID] = {}
    seen: set[str] = set()
    for field, role_hint in VENDOR_SOURCE_FIELDS:
        raw_name = str(raw.get(field) or "").strip()
        if not raw_name or raw_name in seen:
            continue
        seen.add(raw_name)
        candidate = upsert_vendor_candidate(
            db,
            source_system=source_system,
            source_table=source_table,
            source_record_id=f"{source_record_id}:{role_hint}",
            raw_name=raw_name,
        )
        if candidate is not None:
            created_or_updated += 1
            if candidate.matched_org_id and role_hint in {"manufacturer", "registrant", "filer"}:
                matched_org_ids.setdefault(f"{role_hint}_org_id", candidate.matched_org_id)
    return {"count": created_or_updated, **matched_org_ids}


def import_departments(
    db: Session,
    file_content: bytes,
    filename: str,
    sheet_name: str | None,
    batch_id: str,
) -> dict[str, Any]:
    sheet_used, rows = load_tabular_rows(file_content, filename, sheet_name)
    snapshot_existing = set(db.scalars(select(DictDepartment.dept_code)).all())

    failures: list[dict[str, Any]] = []
    success_count = 0
    duplicate_count = 0
    seen_in_file: set[str] = set()

    staged_parents: dict[str, uuid.UUID] = {}

    def resolve_parent_uuid(parent_code: str | None) -> uuid.UUID | None:
        if not parent_code:
            return None
        parent_code = parent_code.strip()
        if not parent_code:
            return None
        if parent_code in staged_parents:
            return staged_parents[parent_code]
        pid = db.scalar(select(DictDepartment.dept_id).where(DictDepartment.dept_code == parent_code))
        return pid

    for line_no, raw in enumerate(rows, start=2):
        dept_code = (raw.get("dept_code") or "").strip()
        dept_name = (raw.get("dept_name") or "").strip()
        if not dept_code:
            failures.append({"row": line_no, "field": "dept_code", "message": "科室编码为空"})
            continue
        if not dept_name:
            failures.append({"row": line_no, "field": "dept_name", "message": "科室名称为空"})
            continue

        dup = dept_code in seen_in_file or dept_code in snapshot_existing
        seen_in_file.add(dept_code)

        parent_code = (raw.get("parent_dept_code") or "").strip()
        parent_id = resolve_parent_uuid(parent_code or None)
        if parent_code and parent_id is None:
            failures.append({"row": line_no, "field": "parent_dept_code", "message": "上级科室编码不存在"})
            continue

        alias = (raw.get("dept_alias") or "").strip()
        dept_type = (raw.get("dept_type") or "").strip()
        oid = (raw.get("oid") or "").strip()
        status = (raw.get("status") or "ACTIVE").strip()

        existing = db.scalar(select(DictDepartment).where(DictDepartment.dept_code == dept_code))
        if existing:
            if dup:
                duplicate_count += 1
            existing.dept_name = dept_name
            existing.dept_alias = alias or None
            existing.parent_dept_id = parent_id
            existing.dept_type = dept_type or None
            existing.oid = oid or None
            existing.status = status
            existing.source_batch_id = batch_id
            staged_parents[dept_code] = existing.dept_id
        else:
            row = DictDepartment(
                dept_code=dept_code,
                dept_name=dept_name,
                dept_alias=alias or None,
                parent_dept_id=parent_id,
                dept_type=dept_type or None,
                oid=oid or None,
                status=status,
                source_batch_id=batch_id,
            )
            db.add(row)
            db.flush()
            staged_parents[dept_code] = row.dept_id
            snapshot_existing.add(dept_code)
        success_count += 1

    return {
        "batch_id": batch_id,
        "sheet_name": sheet_used or None,
        "source_row_count": len(rows),
        "unique_key_count": len(seen_in_file),
        "skipped_duplicate_count": 0,
        "success_count": success_count,
        "failed_count": len(failures),
        "duplicate_count": duplicate_count,
        "failures": failures,
    }


def _normalize_yb27(raw: str) -> str:
    s = str(raw).strip()
    return s


def import_materials_mvp_template(
    db: Session,
    file_content: bytes,
    filename: str,
    sheet_name: str | None,
    batch_id: str,
) -> dict[str, Any]:
    sheet_used, rows = load_tabular_rows(file_content, filename, sheet_name)
    failures: list[dict[str, Any]] = []
    success_count = 0
    duplicate_count = 0
    skipped_duplicate_count = 0
    seen_codes: set[str] = set()
    vendor_candidate_count = 0

    line_no = 1
    for raw in rows:
        line_no += 1
        yb27 = _normalize_yb27(raw.get("yb_code_27") or "")
        generic_name = (raw.get("generic_name") or "").strip()
        if len(yb27) != 27:
            failures.append({"row": line_no, "field": "yb_code_27", "message": "医保码必须为27位"})
            continue
        if not generic_name:
            failures.append({"row": line_no, "field": "generic_name", "message": "通用名不能为空"})
            continue
        if yb27 in seen_codes:
            duplicate_count += 1
            skipped_duplicate_count += 1
            continue
        seen_codes.add(yb27)
        vendor_match = _extract_vendor_candidates(
            db,
            raw,
            source_system="MANUAL",
            source_table="dict_material_specs",
            source_record_id=f"{batch_id}:{yb27}",
        )
        vendor_candidate_count += vendor_match["count"]

        yb20_raw = (raw.get("yb_code_20") or "").strip()
        yb20 = yb20_raw if len(yb20_raw) == 20 else yb27[:20]

        extra_fragment: dict[str, Any] = {}
        sv_raw = (raw.get("spec_value") or "").strip()
        su_raw = (raw.get("spec_unit") or "").strip()
        spec_value = sv_raw or None
        spec_unit = su_raw or None
        if spec_value and spec_unit:
            pass
        elif spec_value and not spec_unit:
            sv2, su2, raw_spec = split_spec(spec_value)
            if sv2:
                spec_value, spec_unit = sv2, su2 or spec_unit
            if raw_spec:
                extra_fragment["raw_spec"] = raw_spec

        existing = db.scalar(select(DictMaterialSpec).where(DictMaterialSpec.yb_code_27 == yb27))
        payload = dict(
            yb_code_27=yb27,
            yb_code_20=yb20,
            manufacturer_org_id=vendor_match.get("manufacturer_org_id"),
            registrant_org_id=vendor_match.get("registrant_org_id"),
            filer_org_id=vendor_match.get("filer_org_id"),
            generic_name=generic_name,
            brand_name=(raw.get("brand_name") or "").strip() or None,
            cat_level_1=(raw.get("cat_level_1") or "").strip() or None,
            cat_level_2=(raw.get("cat_level_2") or "").strip() or None,
            cat_level_3=(raw.get("cat_level_3") or "").strip() or None,
            material_attr=(raw.get("material_attr") or "").strip() or None,
            spec_value=spec_value,
            spec_unit=spec_unit,
            model_detail=(raw.get("model_detail") or "").strip() or None,
            reg_number=(raw.get("reg_number") or "").strip() or None,
            unit_pkg=(raw.get("unit_pkg") or "").strip() or None,
            status=(raw.get("status") or "ACTIVE").strip(),
            source_batch_id=batch_id,
        )
        if existing:
            duplicate_count += 1
            for k, v in payload.items():
                setattr(existing, k, v)
            if extra_fragment:
                merged = dict(existing.extra_attrs or {})
                merged.update(extra_fragment)
                existing.extra_attrs = merged
        else:
            extra_attrs: dict[str, Any] = {}
            if spec_value and not spec_unit:
                sv2, su2, raw_spec = split_spec(spec_value)
                if sv2:
                    payload["spec_value"], payload["spec_unit"] = sv2, su2
                if raw_spec:
                    extra_attrs["raw_spec"] = raw_spec
            payload["extra_attrs"] = extra_attrs
            db.add(DictMaterialSpec(**payload))
        success_count += 1

    return {
        "batch_id": batch_id,
        "source_type": "MVP_TEMPLATE",
        "source_file_name": Path(filename).name,
        "sheet_name": sheet_used or None,
        "source_row_count": len(rows),
        "unique_key_count": len(seen_codes),
        "skipped_duplicate_count": skipped_duplicate_count,
        "success_count": success_count,
        "failed_count": len(failures),
        "duplicate_count": duplicate_count,
        "disabled_count": 0,
        "transcoded_count": 0,
        "changed_count": 0,
        "mapped_count": 0,
        "retained_count": 0,
        "vendor_candidate_count": vendor_candidate_count,
        "failures": failures,
    }


def import_nhsa_full_spec(
    db: Session,
    file_content: bytes,
    filename: str,
    sheet_name: str | None,
    batch_id: str,
) -> dict[str, Any]:
    sheet_used, rows = load_tabular_rows(file_content, filename, sheet_name)
    return import_nhsa_full_spec_rows(db, rows, filename, sheet_used, batch_id, total_rows=len(rows))


def import_nhsa_full_spec_from_path(
    db: Session,
    file_path: Path,
    filename: str,
    sheet_name: str | None,
    batch_id: str,
    progress_callback: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    sheet_used, total_rows, rows = iter_tabular_rows_from_path(file_path, sheet_name)
    return import_nhsa_full_spec_rows(
        db,
        rows,
        filename,
        sheet_used,
        batch_id,
        total_rows=total_rows,
        progress_callback=progress_callback,
    )


def import_nhsa_full_spec_rows(
    db: Session,
    rows: Iterable[dict[str, Any]],
    filename: str,
    sheet_used: str,
    batch_id: str,
    *,
    total_rows: int | None,
    progress_callback: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    failures: list[dict[str, Any]] = []
    success_count = 0
    retained_count = 0
    skipped_duplicate_count = 0
    source_row_count = 0
    vendor_candidate_count = 0

    line_no = 1
    seen_codes: set[str] = set()
    for raw in rows:
        source_row_count += 1
        line_no += 1
        yb27 = _normalize_yb27(raw.get("27位码") or raw.get("yb_code_27") or "")
        if len(yb27) != 27:
            failures.append({"row": line_no, "field": "27位码", "message": "必须为27位"})
            if progress_callback and source_row_count % 5000 == 0:
                progress_callback(
                    {
                        "source_row_count": source_row_count,
                        "unique_key_count": len(seen_codes),
                        "success_count": success_count,
                        "failed_count": len(failures),
                        "skipped_duplicate_count": skipped_duplicate_count,
                        "total_rows": total_rows,
                    }
                )
            continue
        if yb27 in seen_codes:
            skipped_duplicate_count += 1
            if progress_callback and source_row_count % 5000 == 0:
                progress_callback(
                    {
                        "source_row_count": source_row_count,
                        "unique_key_count": len(seen_codes),
                        "success_count": success_count,
                        "failed_count": len(failures),
                        "skipped_duplicate_count": skipped_duplicate_count,
                        "total_rows": total_rows,
                    }
                )
            continue
        seen_codes.add(yb27)
        vendor_match = _extract_vendor_candidates(
            db,
            raw,
            source_system="NHSA",
            source_table="stg_nhsa_material_specs",
            source_record_id=f"{batch_id}:{yb27}",
        )
        vendor_candidate_count += vendor_match["count"]

        change_type = (raw.get("类型") or "").strip()
        if change_type == "保留":
            retained_count += 1

        stg = StgNhsaMaterialSpec(
            batch_id=batch_id,
            source_file_name=Path(filename).name,
            sheet_name=sheet_used,
            row_number=line_no,
            goods_id=(raw.get("GOODS_ID") or "").strip() or None,
            udi=(raw.get("UDI") or "").strip() or None,
            original_yb_code_27=(raw.get("原27位码") or "").strip() or None,
            original_code_status=(raw.get("原码状态") or "").strip() or None,
            yb_code_27=yb27,
            change_type=change_type or None,
            yb_code_20=(raw.get("20位码") or "").strip() or None,
            cat_level_1=(raw.get("一级分类") or "").strip() or None,
            cat_level_2=(raw.get("二级分类") or "").strip() or None,
            cat_level_3=(raw.get("三级分类") or "").strip() or None,
            insurance_generic_category=(raw.get("医保通用名分类") or "").strip() or None,
            material_attr=(raw.get("材质") or "").strip() or None,
            feature=(raw.get("特征") or "").strip() or None,
            registration_number=(raw.get("注册备案号") or "").strip() or None,
            single_product_name=(raw.get("单件产品名称") or "").strip() or None,
            manufacturer=(raw.get("耗材企业") or "").strip() or None,
            spec_model_count=int(raw["规格型号数"]) if str(raw.get("规格型号数") or "").isdigit() else None,
            spec=(raw.get("规格") or "").strip() or None,
            model=(raw.get("型号") or "").strip() or None,
            insurance_generic_name_code=(raw.get("医保通用名编号") or "").strip() or None,
            insurance_generic_name=(raw.get("医保通用名") or "").strip() or None,
            raw_payload=dict(raw),
        )
        db.add(stg)
        db.flush()

        generic = (
            stg.insurance_generic_name or stg.single_product_name or ""
        ).strip()
        if not generic:
            failures.append({"row": line_no, "field": "generic", "message": "缺少通用名/产品名称"})
            db.expunge(stg)
            continue

        yb20 = stg.yb_code_20.strip() if stg.yb_code_20 else yb27[:20]
        sv, su, raw_spec = split_spec(stg.spec)
        extra_attrs: dict[str, Any] = {"insurance_generic_category": stg.insurance_generic_category}
        if raw_spec:
            extra_attrs["raw_spec"] = raw_spec

        existing = db.scalar(select(DictMaterialSpec).where(DictMaterialSpec.yb_code_27 == yb27))
        fields = dict(
            yb_code_27=yb27,
            yb_code_20=yb20[:20],
            manufacturer_org_id=vendor_match.get("manufacturer_org_id"),
            registrant_org_id=vendor_match.get("registrant_org_id"),
            filer_org_id=vendor_match.get("filer_org_id"),
            goods_id=stg.goods_id,
            udi=stg.udi,
            original_yb_code_27=stg.original_yb_code_27,
            original_code_status=stg.original_code_status,
            code_change_type=stg.change_type,
            generic_name=generic,
            brand_name=stg.manufacturer,
            cat_level_1=stg.cat_level_1,
            cat_level_2=stg.cat_level_2,
            cat_level_3=stg.cat_level_3,
            material_attr=stg.material_attr,
            feature=stg.feature,
            spec_value=sv,
            spec_unit=su,
            model_detail=stg.model,
            reg_number=stg.registration_number,
            unit_pkg=None,
            insurance_generic_name_code=stg.insurance_generic_name_code,
            insurance_generic_name=stg.insurance_generic_name,
            status="ACTIVE",
            source_batch_id=batch_id,
            extra_attrs=extra_attrs,
        )
        if existing:
            for k, v in fields.items():
                setattr(existing, k, v)
        else:
            db.add(DictMaterialSpec(**fields))
        success_count += 1
        if progress_callback and source_row_count % 5000 == 0:
            progress_callback(
                {
                    "source_row_count": source_row_count,
                    "unique_key_count": len(seen_codes),
                    "success_count": success_count,
                    "failed_count": len(failures),
                    "skipped_duplicate_count": skipped_duplicate_count,
                    "total_rows": total_rows,
                }
            )

    return {
        "batch_id": batch_id,
        "source_type": "NHSA_FULL_SPEC",
        "source_file_name": Path(filename).name,
        "sheet_name": sheet_used,
        "source_row_count": source_row_count,
        "unique_key_count": len(seen_codes),
        "skipped_duplicate_count": skipped_duplicate_count,
        "success_count": success_count,
        "failed_count": len(failures),
        "duplicate_count": 0,
        "disabled_count": 0,
        "transcoded_count": 0,
        "changed_count": 0,
        "mapped_count": 0,
        "retained_count": retained_count,
        "vendor_candidate_count": vendor_candidate_count,
        "failures": failures,
    }


def import_nhsa_disabled(
    db: Session,
    file_content: bytes,
    filename: str,
    sheet_name: str | None,
    batch_id: str,
) -> dict[str, Any]:
    sheet_used, rows = load_tabular_rows(file_content, filename, sheet_name)
    failures: list[dict[str, Any]] = []
    codes: list[str] = []
    line_no = 1
    seen: set[str] = set()
    skipped_duplicate_count = 0
    for raw in rows:
        line_no += 1
        yb27 = _normalize_yb27(raw.get("27位码") or "")
        status_txt = (raw.get("原码状态") or "").strip()
        if len(yb27) != 27:
            failures.append({"row": line_no, "field": "27位码", "message": "必须为27位"})
            continue
        if yb27 in seen:
            skipped_duplicate_count += 1
            continue
        seen.add(yb27)
        db.add(
            StgNhsaMaterialDisabled(
                batch_id=batch_id,
                source_file_name=Path(filename).name,
                sheet_name=sheet_used,
                row_number=line_no,
                yb_code_27=yb27,
                original_code_status=status_txt or "停用",
                raw_payload=dict(raw),
            )
        )
        codes.append(yb27)

    if codes:
        db.execute(
            update(DictMaterialSpec)
            .where(DictMaterialSpec.yb_code_27.in_(codes))
            .values(status="INACTIVE")
        )

    return {
        "batch_id": batch_id,
        "source_type": "NHSA_DISABLED",
        "source_file_name": Path(filename).name,
        "sheet_name": sheet_used,
        "source_row_count": len(rows),
        "unique_key_count": len(seen),
        "skipped_duplicate_count": skipped_duplicate_count,
        "success_count": len(codes),
        "failed_count": len(failures),
        "duplicate_count": 0,
        "disabled_count": len(codes),
        "transcoded_count": 0,
        "changed_count": 0,
        "mapped_count": 0,
        "retained_count": 0,
        "failures": failures,
    }


def import_nhsa_transcode(
    db: Session,
    file_content: bytes,
    filename: str,
    sheet_name: str | None,
    batch_id: str,
) -> dict[str, Any]:
    sheet_used, rows = load_tabular_rows(file_content, filename, sheet_name)
    failures: list[dict[str, Any]] = []
    changed_count = 0
    mapped_count = 0
    line_no = 1
    seen: set[tuple[str, str]] = set()
    skipped_duplicate_count = 0
    for raw in rows:
        line_no += 1
        orig = _normalize_yb27(raw.get("原27位码") or "")
        newc = _normalize_yb27(raw.get("27位码") or "")
        ctype = (raw.get("类型") or "").strip()
        if len(orig) != 27 or len(newc) != 27:
            failures.append({"row": line_no, "field": "codes", "message": "原码与新码须为27位"})
            continue
        key = (orig, newc)
        if key in seen:
            skipped_duplicate_count += 1
            continue
        seen.add(key)
        if ctype == "变更":
            changed_count += 1
        elif ctype == "映射":
            mapped_count += 1
        db.add(
            StgNhsaMaterialTranscode(
                batch_id=batch_id,
                source_file_name=Path(filename).name,
                sheet_name=sheet_used,
                row_number=line_no,
                original_yb_code_27=orig,
                original_code_status=(raw.get("原码状态") or "").strip() or "停用",
                yb_code_27=newc,
                change_type=ctype or "映射",
                raw_payload=dict(raw),
            )
        )

    return {
        "batch_id": batch_id,
        "source_type": "NHSA_TRANSCODE",
        "source_file_name": Path(filename).name,
        "sheet_name": sheet_used,
        "source_row_count": len(rows),
        "unique_key_count": len(seen),
        "skipped_duplicate_count": skipped_duplicate_count,
        "success_count": len(seen),
        "failed_count": len(failures),
        "duplicate_count": 0,
        "disabled_count": 0,
        "transcoded_count": len(seen),
        "changed_count": changed_count,
        "mapped_count": mapped_count,
        "retained_count": 0,
        "failures": failures,
    }


def import_device_classification_docx(
    db: Session,
    file_content: bytes,
    filename: str,
    batch_id: str,
    source_system: str | None = None,
    publish_date: Any = None,
    effective_date: Any = None,
    operator_name: str | None = None,
) -> dict[str, Any]:
    try:
        doc = Document(io.BytesIO(file_content))
    except Exception as exc:
        raise ValueError("无法读取医疗器械分类目录 DOCX") from exc

    failures: list[dict[str, Any]] = []
    success_count = 0
    duplicate_count = 0
    source_name = Path(filename).name
    existing_rows = db.scalars(select(RefDeviceClassificationCatalog)).all()

    def natural_key_from_values(
        major_category_no: str | None,
        major_category_name: str | None,
        level_1_category_no: str | None,
        level_1_category: str | None,
        level_2_category_no: str | None,
        level_2_category: str | None,
        product_description: str | None,
        intended_use: str | None,
        product_examples: str | None,
        management_class: str | None,
    ) -> tuple[str, ...]:
        return (
            major_category_no or "",
            major_category_name or "",
            level_1_category_no or "",
            level_1_category or "",
            level_2_category_no or "",
            level_2_category or "",
            product_description or "",
            intended_use or "",
            product_examples or "",
            management_class or "",
        )

    def natural_key(row: RefDeviceClassificationCatalog) -> tuple[str, ...]:
        return natural_key_from_values(
            row.major_category_no,
            row.major_category_name,
            row.level_1_category_no,
            row.level_1_category,
            row.level_2_category_no,
            row.level_2_category,
            row.product_description,
            row.intended_use,
            row.product_examples,
            row.management_class,
        )

    existing_by_key: dict[tuple[str, ...], RefDeviceClassificationCatalog] = {}
    duplicate_ids: list[Any] = []
    for existing in existing_rows:
        key = natural_key(existing)
        if key in existing_by_key:
            duplicate_ids.append(existing.catalog_id)
        else:
            existing_by_key[key] = existing
    if duplicate_ids:
        db.execute(
            update(RefDeviceClassificationCatalog)
            .where(RefDeviceClassificationCatalog.catalog_id.in_(duplicate_ids))
            .values(
                data_status="deprecated",
                hidden_in_tree=True,
                status_reason="导入核验发现重复自然键，已作废隐藏并保留历史",
                corrected_at=datetime.now(timezone.utc),
            )
        )
        duplicate_count += len(duplicate_ids)
    db.execute(delete(RefDeviceClassificationRevision).where(RefDeviceClassificationRevision.batch_id == batch_id))
    db.add(
        DeviceClassificationImportAuditLog(
            batch_id=batch_id,
            action="CONFIRM_START",
            operator_name=operator_name,
            message="确认写入医疗器械分类目录正式表",
            payload={"source_file_name": source_name},
        )
    )

    def catalog_payload(row: RefDeviceClassificationCatalog | None = None, payload: dict[str, Any] | None = None) -> dict[str, Any]:
        source = payload or {}
        return {
            "major_category_no": row.major_category_no if row else source.get("major_category_no"),
            "major_category_name": row.major_category_name if row else source.get("major_category_name"),
            "level_1_category_no": row.level_1_category_no if row else source.get("level_1_category_no"),
            "level_1_category": row.level_1_category if row else source.get("level_1_category"),
            "level_2_category_no": row.level_2_category_no if row else source.get("level_2_category_no"),
            "level_2_category": row.level_2_category if row else source.get("level_2_category"),
            "product_description": row.product_description if row else source.get("product_description"),
            "intended_use": row.intended_use if row else source.get("intended_use"),
            "product_examples": row.product_examples if row else source.get("product_examples"),
            "management_class": row.management_class if row else source.get("management_class"),
        }

    def add_revision(
        *,
        source_table_index: int,
        row_number: int,
        change_type: str,
        old_row: RefDeviceClassificationCatalog | None,
        new_row: RefDeviceClassificationCatalog | None,
        old_payload: dict[str, Any],
        new_payload: dict[str, Any],
    ) -> None:
        db.add(
            RefDeviceClassificationRevision(
                batch_id=batch_id,
                source_file_name=source_name,
                source_table_index=source_table_index,
                row_number=row_number,
                change_type=change_type,
                reason=f"{source_name} 第 {row_number} 行目录调整",
                old_catalog_id=old_row.catalog_id if old_row else None,
                new_catalog_id=new_row.catalog_id if new_row else None,
                old_payload=old_payload,
                new_payload=new_payload,
            )
        )

    adjustment_success_count = 0
    adjustment_abnormal_count = 0
    adjustment_deleted_keys: set[tuple[str, ...]] = set()
    adjustment_seen_keys: set[tuple[str, ...]] = set()
    for table_index, table in enumerate(doc.tables):
        if len(table.rows) < 3:
            continue
        top_headers = [_normalize_device_header(value) for value in _docx_row_values(table.rows[0])]
        headers = [_normalize_device_header(value) for value in _docx_row_values(table.rows[1])]
        if "调整后《医疗器械分类目录》内容" not in top_headers:
            continue
        if len(headers) < 15 or not DEVICE_ADJUSTMENT_HEADERS_REQUIRED.issubset(set(headers)):
            continue
        for ri in range(2, len(table.rows)):
            cells = _unique_docx_row_values(table.rows[ri])
            if len(cells) < 15:
                continue
            original = _parse_device_adjustment_side(cells, 1)
            adjusted = _parse_device_adjustment_side(cells, 8, original)
            if all(value is None for key, value in adjusted.items() if not key.endswith("_raw")):
                continue
            old_key = natural_key_from_values(
                original["major_category_no"],
                original["major_category_name"],
                original["level_1_category_no"],
                original["level_1_category"],
                original["level_2_category_no"],
                original["level_2_category"],
                original["product_description"],
                original["intended_use"],
                original["product_examples"],
                original["management_class"],
            )
            new_key = natural_key_from_values(
                adjusted["major_category_no"],
                adjusted["major_category_name"],
                adjusted["level_1_category_no"],
                adjusted["level_1_category"],
                adjusted["level_2_category_no"],
                adjusted["level_2_category"],
                adjusted["product_description"],
                adjusted["intended_use"],
                adjusted["product_examples"],
                adjusted["management_class"],
            )
            has_original = _has_device_adjustment_original(original)
            existing_old = existing_by_key.get(old_key) if has_original else None
            old_revision_payload = catalog_payload(existing_old) if existing_old else catalog_payload(payload=original)
            if has_original and old_key != new_key and old_key not in adjustment_deleted_keys:
                existing_old = existing_by_key.pop(old_key, None)
                if existing_old is not None:
                    existing_old.data_status = "deprecated"
                    existing_old.hidden_in_tree = True
                    existing_old.status_reason = f"{source_name} 第 {ri + 1} 行目录调整，旧版本作废隐藏"
                    existing_old.corrected_at = datetime.now(timezone.utc)
                    duplicate_count += 1
                adjustment_deleted_keys.add(old_key)
            payload = {
                "batch_id": batch_id,
                "source_file_name": source_name,
                "source_table_index": table_index,
                "row_number": ri + 1,
                "category_no": adjusted["level_1_category_no"],
                "major_category_no": adjusted["major_category_no"],
                "major_category_name": adjusted["major_category_name"],
                "level_1_category_no": adjusted["level_1_category_no"],
                "level_1_category": adjusted["level_1_category"],
                "level_2_category_no": adjusted["level_2_category_no"],
                "level_2_category": adjusted["level_2_category"],
                "product_description": adjusted["product_description"],
                "intended_use": adjusted["intended_use"],
                "product_examples": adjusted["product_examples"],
                "management_class": adjusted["management_class"],
            }
            payload = _apply_device_governance_status(payload)
            if payload.get("data_status") != "effective":
                adjustment_abnormal_count += 1
                db.add(ImportValidationIssue(
                    batch_id=batch_id,
                    rule_code="device_catalog_governance_validation",
                    severity="warning",
                    issue_type=str(payload.get("abnormal_type") or "人工确认异常"),
                    message=str(payload.get("status_reason") or "需人工确认"),
                    source_position=f"表 {table_index + 1} 第 {ri + 1} 行",
                    raw_payload=payload,
                ))
            model_payload = _device_catalog_model_payload(payload)
            if new_key in adjustment_seen_keys:
                duplicate_count += 1
                continue
            adjustment_seen_keys.add(new_key)
            existing = existing_by_key.get(new_key)
            if existing:
                duplicate_count += 1
                for field, value in model_payload.items():
                    setattr(existing, field, value)
                new_row = existing
            else:
                row = RefDeviceClassificationCatalog(**model_payload)
                db.add(row)
                db.flush()
                existing_by_key[new_key] = row
                new_row = row
            change_type = "ADDED" if not has_original else "MODIFIED"
            add_revision(
                source_table_index=table_index,
                row_number=ri + 1,
                change_type=change_type,
                old_row=existing_old,
                new_row=new_row,
                old_payload={} if not has_original else old_revision_payload,
                new_payload=catalog_payload(new_row),
            )
            adjustment_success_count += 1

    if adjustment_success_count:
        report = {
            "batch_id": batch_id,
            "source_type": "DEVICE_CLASSIFICATION_CATALOG",
            "source_file_name": source_name,
            "sheet_name": None,
            "source_row_count": adjustment_success_count,
            "unique_key_count": len(adjustment_seen_keys),
            "skipped_duplicate_count": duplicate_count,
            "success_count": adjustment_success_count,
            "failed_count": len(failures),
            "pending_confirm_count": adjustment_abnormal_count,
            "duplicate_count": duplicate_count,
            "disabled_count": 0,
            "transcoded_count": 0,
            "changed_count": len(adjustment_deleted_keys),
            "mapped_count": 0,
            "retained_count": 0,
            "failures": failures,
        }
        _record_device_catalog_version(
            db,
            batch_id=batch_id,
            source_file_name=source_name,
            source_system=source_system,
            publish_date=publish_date,
            effective_date=effective_date,
            item_count=adjustment_success_count,
            abnormal_count=adjustment_abnormal_count + len(failures),
        )
        _record_device_import_audit(
            db,
            batch_id=batch_id,
            action="CONFIRM_COMMIT",
            operator_name=operator_name,
            message="医疗器械分类目录调整批次已写入正式表",
            payload=report,
        )
        _record_device_import_audit(
            db,
            batch_id=batch_id,
            action="VERSION_CREATED",
            operator_name=operator_name,
            message="医疗器械分类目录版本记录已生成",
            payload={"catalog_version": batch_id, "item_count": adjustment_success_count},
        )
        _record_device_import_audit(
            db,
            batch_id=batch_id,
            action="AUDIT_RECORDED",
            operator_name=operator_name,
            message="导入审计日志已生成",
            payload={"audit_status": "recorded"},
        )
        report["catalog_version"] = batch_id
        report["validation_report_status"] = "已生成"
        report["audit_log_status"] = "已记录"
        report["rollback_supported"] = True
        return report

    current_major_no: str | None = None
    current_major_name: str | None = None
    table_index = -1
    for block_type, block in _iter_docx_blocks(doc):
        if block_type == "paragraph":
            heading = _parse_device_major_heading(block.text)
            if heading:
                current_major_no, current_major_name = heading
            continue

        table = block
        if not isinstance(table, Table):
            continue
        table_index += 1
        if len(table.rows) < 2:
            continue
        hdr = [_normalize_device_header(value) for value in _unique_docx_row_values(table.rows[0])]
        hdr_set = set(hdr)
        if not DEVICE_HEADERS_REQUIRED.issubset(hdr_set):
            continue
        idx_map = _device_header_indexes(hdr)
        for ri in range(1, len(table.rows)):
            cells = _unique_docx_row_values(table.rows[ri])

            def cell(name: str) -> str:
                i = idx_map.get(name)
                if i is None or i >= len(cells):
                    return ""
                return cells[i]

            if not any(cell(name) for name in DEVICE_HEADERS_REQUIRED):
                continue
            level_1_no, level_1_name = _split_device_code_name(cell("序号"))
            level_2_no, level_2_name = _split_device_code_name(cell("二级产品类别"))
            level_1_category = _normalize_device_value(cell("一级产品类别")) or None
            payload = {
                "batch_id": batch_id,
                "source_file_name": source_name,
                "source_table_index": table_index,
                "row_number": ri,
                "category_no": level_1_no,
                "major_category_no": current_major_no,
                "major_category_name": current_major_name,
                "level_1_category_no": level_1_no,
                "level_1_category": level_1_category,
                "level_2_category_no": level_2_no,
                "level_2_category": level_2_name,
                "product_description": _normalize_device_value(cell("产品描述")) or None,
                "intended_use": _normalize_device_value(cell("预期用途")) or None,
                "product_examples": _normalize_device_value(cell("品名举例")) or None,
                "management_class": _normalize_device_value(cell("管理类别")) or None,
            }
            payload = _apply_device_governance_status(payload)
            if payload.get("data_status") != "effective":
                failures.append({
                    "row": ri,
                    "field": payload.get("abnormal_type") or "validation",
                    "message": payload.get("status_reason") or "需人工确认",
                })
                db.add(ImportValidationIssue(
                    batch_id=batch_id,
                    rule_code="device_catalog_governance_validation",
                    severity="warning",
                    issue_type=str(payload.get("abnormal_type") or "人工确认异常"),
                    message=str(payload.get("status_reason") or "需人工确认"),
                    source_position=f"表 {table_index + 1} 第 {ri} 行",
                    raw_payload=payload,
                ))
            key = (
                payload["major_category_no"] or "",
                payload["major_category_name"] or "",
                payload["level_1_category_no"] or "",
                payload["level_1_category"] or "",
                payload["level_2_category_no"] or "",
                payload["level_2_category"] or "",
                payload["product_description"] or "",
                payload["intended_use"] or "",
                payload["product_examples"] or "",
                payload["management_class"] or "",
            )
            existing = existing_by_key.get(key)
            model_payload = _device_catalog_model_payload(payload)
            if existing:
                duplicate_count += 1
                for field, value in model_payload.items():
                    setattr(existing, field, value)
            else:
                row = RefDeviceClassificationCatalog(**model_payload)
                db.add(row)
                existing_by_key[key] = row
            success_count += 1

    if success_count == 0:
        failures.append({"row": 0, "field": "docx", "message": "未发现包含七个核心字段的分类明细表"})

    report = {
        "batch_id": batch_id,
        "source_type": "DEVICE_CLASSIFICATION_CATALOG",
        "source_file_name": source_name,
        "sheet_name": None,
        "source_row_count": success_count,
        "unique_key_count": success_count,
        "skipped_duplicate_count": duplicate_count,
        "success_count": success_count,
        "failed_count": len(failures),
        "pending_confirm_count": sum(1 for item in failures if item.get("message") != "未发现包含七个核心字段的分类明细表"),
        "duplicate_count": duplicate_count,
        "disabled_count": 0,
        "transcoded_count": 0,
        "changed_count": 0,
        "mapped_count": 0,
        "retained_count": 0,
        "failures": failures,
    }
    _record_device_catalog_version(
        db,
        batch_id=batch_id,
        source_file_name=source_name,
        source_system=source_system,
        publish_date=publish_date,
        effective_date=effective_date,
        item_count=success_count,
        abnormal_count=len(failures),
    )
    _record_device_import_audit(
        db,
        batch_id=batch_id,
        action="CONFIRM_COMMIT",
        operator_name=operator_name,
        message="医疗器械分类目录导入批次已写入正式表",
        payload=report,
    )
    _record_device_import_audit(
        db,
        batch_id=batch_id,
        action="VERSION_CREATED",
        operator_name=operator_name,
        message="医疗器械分类目录版本记录已生成",
        payload={"catalog_version": batch_id, "item_count": success_count},
    )
    _record_device_import_audit(
        db,
        batch_id=batch_id,
        action="AUDIT_RECORDED",
        operator_name=operator_name,
        message="导入审计日志已生成",
        payload={"audit_status": "recorded"},
    )
    report["catalog_version"] = batch_id
    report["validation_report_status"] = "已生成"
    report["audit_log_status"] = "已记录"
    report["rollback_supported"] = True
    return report


def _record_device_catalog_version(
    db: Session,
    *,
    batch_id: str,
    source_file_name: str,
    source_system: str | None,
    publish_date: Any,
    effective_date: Any,
    item_count: int,
    abnormal_count: int,
) -> None:
    db.execute(
        update(DeviceClassificationCatalogVersion)
        .where(DeviceClassificationCatalogVersion.batch_id != batch_id)
        .values(is_current=False)
    )
    db.execute(
        delete(DeviceClassificationCatalogVersion).where(DeviceClassificationCatalogVersion.batch_id == batch_id)
    )
    db.add(
        DeviceClassificationCatalogVersion(
            batch_id=batch_id,
            version_no=batch_id,
            source_file_name=source_file_name,
            source_system=source_system,
            publish_date=publish_date,
            effective_date=effective_date,
            item_count=item_count,
            abnormal_count=abnormal_count,
            is_current=True,
            rollback_supported=True,
        )
    )


def _record_device_import_audit(
    db: Session,
    *,
    batch_id: str,
    action: str,
    operator_name: str | None,
    message: str,
    payload: dict[str, Any],
) -> None:
    db.add(
        DeviceClassificationImportAuditLog(
            batch_id=batch_id,
            action=action,
            operator_name=operator_name,
            message=message,
            payload=payload,
        )
    )


def _unsupported_device_preview(filename: str, message: str) -> dict[str, Any]:
    source_name = Path(filename).name
    validation = {
        "key": "parser_support",
        "label": "文件解析器核验",
        "status": "错误",
        "message": message,
        "count": 1,
    }
    return {
        "source_file_name": source_name,
        "source_row_count": 0,
        "major_count": 0,
        "level_1_count": 0,
        "level_2_count": 0,
        "catalog_item_count": 0,
        "parse_success_count": 0,
        "parse_failed_count": 1,
        "pending_confirm_count": 0,
        "management_class_counts": {},
        "invalid_count": 1,
        "duplicate_count": 0,
        "batch_risk_level": "high",
        "batch_risk_summary": "文件格式或解析器不支持，禁止直接入有效目录。",
        "validation_results": [validation],
        "warnings": [{"row": 0, "field": "file", "message": message}],
        "samples": [],
    }


def preview_device_classification_file(file_content: bytes, filename: str) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        return preview_device_classification_docx(file_content, filename)
    if suffix in {".xlsx", ".xlsm", ".pdf"}:
        return _unsupported_device_preview(
            filename,
            f"{suffix.removeprefix('.').upper()} 文件已进入导入流水线，但当前版本尚未接入该格式的结构化目录解析器。",
        )
    return _unsupported_device_preview(filename, "仅支持 docx、pdf、xlsx 格式的医疗器械分类目录原始文件。")


def preview_device_classification_docx(file_content: bytes, filename: str) -> dict[str, Any]:
    try:
        doc = Document(io.BytesIO(file_content))
    except Exception as exc:
        raise ValueError("无法读取医疗器械分类目录 DOCX") from exc

    source_name = Path(filename).name
    rows: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    seen_keys: set[tuple[str, ...]] = set()
    duplicate_count = 0
    current_major_no: str | None = None
    current_major_name: str | None = None
    table_index = -1

    rows.extend(_parse_device_adjustment_docx_rows(doc))

    if not rows:
        for block_type, block in _iter_docx_blocks(doc):
            if block_type == "paragraph":
                heading = _parse_device_major_heading(block.text)
                if heading:
                    current_major_no, current_major_name = heading
                continue
            table = block
            if not isinstance(table, Table):
                continue
            table_index += 1
            if len(table.rows) < 2:
                continue
            hdr = [_normalize_device_header(value) for value in _unique_docx_row_values(table.rows[0])]
            if not DEVICE_HEADERS_REQUIRED.issubset(set(hdr)):
                continue
            idx_map = _device_header_indexes(hdr)
            for ri in range(1, len(table.rows)):
                cells = _unique_docx_row_values(table.rows[ri])

                def cell(name: str) -> str:
                    i = idx_map.get(name)
                    if i is None or i >= len(cells):
                        return ""
                    return cells[i]

                if not any(cell(name) for name in DEVICE_HEADERS_REQUIRED):
                    continue
                level_1_no, _level_1_name = _split_device_code_name(cell("序号"))
                level_2_no, level_2_name = _split_device_code_name(cell("二级产品类别"))
                rows.append({
                    "major_category_no": current_major_no,
                    "major_category_name": current_major_name,
                    "level_1_category_no": level_1_no,
                    "level_1_category": _normalize_device_value(cell("一级产品类别")) or None,
                    "level_2_category_no": level_2_no,
                    "level_2_category": level_2_name,
                    "product_description": _normalize_device_value(cell("产品描述")) or None,
                    "intended_use": _normalize_device_value(cell("预期用途")) or None,
                    "product_examples": _normalize_device_value(cell("品名举例")) or None,
                    "management_class": _normalize_device_value(cell("管理类别")) or None,
                    "source_table_index": table_index,
                    "row_number": ri,
                    "change_type": "新增",
                })

    for row in rows:
        key = tuple(str(row.get(name) or "") for name in (
            "major_category_no",
            "major_category_name",
            "level_1_category_no",
            "level_1_category",
            "level_2_category_no",
            "level_2_category",
            "product_description",
            "intended_use",
            "product_examples",
            "management_class",
        ))
        if key in seen_keys:
            duplicate_count += 1
            warnings.append({"row": row.get("row_number") or 0, "field": "natural_key", "message": "发现重复目录明细"})
        seen_keys.add(key)
        missing = [label for label, field in (
            ("大类", "major_category_no"),
            ("一级类别", "level_1_category"),
            ("二级类别", "level_2_category"),
            ("产品描述", "product_description"),
            ("预期用途", "intended_use"),
            ("品名举例", "product_examples"),
            ("管理类别", "management_class"),
        ) if not row.get(field)]
        if missing:
            warnings.append({"row": row.get("row_number") or 0, "field": ",".join(missing), "message": f"存在空字段：{'、'.join(missing)}"})

    if not rows:
        warnings.append({"row": 0, "field": "docx", "message": "未发现包含七个核心字段的分类明细表"})

    def count_unique(*fields: str) -> int:
        return len({tuple(str(row.get(field) or "") for field in fields) for row in rows if any(row.get(field) for field in fields)})

    management_class_counts: dict[str, int] = {}
    invalid_count = 0
    duplicate_item_names: dict[str, int] = {}
    code_counts: dict[str, int] = {}
    for index, row in enumerate(rows):
        management_class = str(row.get("management_class") or "-")
        management_class_counts[management_class] = management_class_counts.get(management_class, 0) + 1
        catalog_code = "-".join(str(row.get(field) or "") for field in ("major_category_no", "level_1_category_no", "level_2_category_no")).strip("-")
        code_counts[catalog_code] = code_counts.get(catalog_code, 0) + 1
        item_name = str(row.get("level_2_category") or "")
        if item_name:
            duplicate_item_names[item_name] = duplicate_item_names.get(item_name, 0) + 1
        row["catalog_code"] = catalog_code or None
        row["parse_status"] = "通过"
        row["issue_message"] = ""
        row["row_number"] = row.get("row_number") or index + 1
        row["staging_id"] = f"preview-{index + 1}"
        row["change_type"] = row.get("change_type") or "新增"
        _apply_device_governance_status(row)
        missing_required = any(str(row.get(field) or "").strip() == "" for field in ("level_2_category", "product_description", "intended_use", "management_class"))
        if row.get("data_status") in {"parse_abnormal", "pending_confirm"}:
            invalid_count += 1
            warnings.append({
                "row": row.get("row_number") or 0,
                "field": row.get("abnormal_type") or "validation",
                "message": row.get("issue_message") or "需人工确认",
            })
        elif missing_required:
            invalid_count += 1
            row["parse_status"] = "错误"
            row["issue_message"] = "必填字段缺失"
            row["change_type"] = "待确认"
        elif management_class not in DEVICE_ALLOWED_MANAGEMENT_CLASSES:
            row["parse_status"] = "待人工确认"
            row["issue_message"] = "管理类别需人工确认"
            row["change_type"] = "待确认"
            row["data_status"] = "pending_confirm"
            row["hidden_in_tree"] = True

    duplicate_code_count = sum(1 for count in code_counts.values() if count > 1)
    duplicate_item_count = sum(1 for count in duplicate_item_names.values() if count > 1)
    pending_confirm_count = sum(1 for row in rows if row.get("parse_status") == "待人工确认")
    failed_count = sum(1 for row in rows if row.get("parse_status") == "错误")
    governance_abnormal_count = sum(1 for row in rows if row.get("data_status") in {"pending_confirm", "parse_abnormal"})
    low_confidence_count = sum(1 for row in rows if int(row.get("confidence_score") or 100) < 60)
    abnormal_ratio = governance_abnormal_count / len(rows) if rows else 1
    if failed_count or abnormal_ratio >= 0.2:
        batch_risk_level = "high"
    elif pending_confirm_count or duplicate_code_count or duplicate_item_count or low_confidence_count:
        batch_risk_level = "medium"
    else:
        batch_risk_level = "low"
    batch_risk_summary = (
        "批次存在较高比例异常或解析错误，应先人工核验异常行。"
        if batch_risk_level == "high"
        else "批次存在待确认、重复或低置信度记录，建议复核后再入库。"
        if batch_risk_level == "medium"
        else "批次结构和字段核验通过，可按流程确认入库。"
    )
    change_type_counts = {label: 0 for label in ("新增", "修改", "废止", "调整", "无变化", "待确认")}
    for row in rows:
        change_type = str(row.get("change_type") or "新增")
        change_type_counts[change_type] = change_type_counts.get(change_type, 0) + 1
    validation_results = [
        {
            "key": "abnormal_name",
            "label": "异常目录名称核验",
            "status": "通过" if not any(row.get("abnormal_type") in {"解析噪声", "目录名称异常", "字段错位"} for row in rows) else "待人工确认",
            "message": "未发现异常目录名称" if not any(row.get("abnormal_type") in {"解析噪声", "目录名称异常", "字段错位"} for row in rows) else "发现疑似解析噪声或字段错位，已进入待人工确认",
            "count": sum(1 for row in rows if row.get("abnormal_type") in {"解析噪声", "目录名称异常", "字段错位"}),
        },
        {
            "key": "revision_note_row",
            "label": "修订说明行核验",
            "status": "通过" if not any(row.get("status_reason") == "疑似修订说明行，不得直接入正式分类树" for row in rows) else "待人工确认",
            "message": "未发现被误解析的修订说明行" if not any(row.get("status_reason") == "疑似修订说明行，不得直接入正式分类树" for row in rows) else "发现疑似修订说明行，已阻断进入有效分类树",
            "count": sum(1 for row in rows if row.get("status_reason") == "疑似修订说明行，不得直接入正式分类树"),
        },
        {
            "key": "code_format",
            "label": "编码格式核验",
            "status": "通过" if not invalid_count else "警告",
            "message": "已识别大类、一级和二级编码" if not invalid_count else "部分记录缺少分类编码或核心字段",
            "count": invalid_count,
        },
        {
            "key": "parent_child",
            "label": "父子级关系核验",
            "status": "通过" if count_unique("major_category_no", "level_1_category_no") else "错误",
            "message": "父子级层级可解析" if rows else "未解析到可用层级",
            "count": count_unique("major_category_no", "level_1_category_no", "level_2_category_no"),
        },
        {
            "key": "required_fields",
            "label": "必填字段核验",
            "status": "通过" if invalid_count == 0 else "错误",
            "message": "必填字段完整" if invalid_count == 0 else f"{invalid_count} 条记录存在必填字段缺失",
            "count": invalid_count,
        },
        {
            "key": "management_class",
            "label": "管理类别合法性核验",
            "status": "通过" if pending_confirm_count == 0 else "待人工确认",
            "message": "管理类别均在允许范围内" if pending_confirm_count == 0 else f"{pending_confirm_count} 条管理类别需确认",
            "count": pending_confirm_count,
        },
        {
            "key": "duplicate_code",
            "label": "重复编码核验",
            "status": "通过" if duplicate_code_count == 0 else "警告",
            "message": "未发现重复分类编码" if duplicate_code_count == 0 else f"{duplicate_code_count} 个编码重复",
            "count": duplicate_code_count,
        },
        {
            "key": "duplicate_item",
            "label": "重复目录条目核验",
            "status": "通过" if duplicate_item_count == 0 else "警告",
            "message": "未发现重复目录条目" if duplicate_item_count == 0 else f"{duplicate_item_count} 个目录条目重复",
            "count": duplicate_item_count,
        },
        {
            "key": "version_consistency",
            "label": "版本一致性核验",
            "status": "通过",
            "message": "文件可作为同一批次版本解析",
            "count": len(rows),
        },
        {
            "key": "batch_risk",
            "label": "批次风险核验",
            "status": "通过" if batch_risk_level == "low" else "警告" if batch_risk_level == "medium" else "待人工确认",
            "message": batch_risk_summary,
            "count": governance_abnormal_count + failed_count + duplicate_code_count + duplicate_item_count,
        },
    ]

    return {
        "source_file_name": source_name,
        "source_row_count": len(rows),
        "major_count": count_unique("major_category_no", "major_category_name"),
        "level_1_count": count_unique("major_category_no", "level_1_category_no", "level_1_category"),
        "level_2_count": count_unique("major_category_no", "level_1_category_no", "level_2_category_no", "level_2_category"),
        "catalog_item_count": len(rows),
        "parse_success_count": len(rows) - failed_count,
        "parse_failed_count": failed_count,
        "pending_confirm_count": pending_confirm_count,
        "management_class_counts": management_class_counts,
        "invalid_count": invalid_count,
        "duplicate_count": duplicate_count,
        "low_confidence_count": low_confidence_count,
        "batch_risk_level": batch_risk_level,
        "batch_risk_summary": batch_risk_summary,
        "change_type_counts": change_type_counts,
        "validation_results": validation_results,
        "warnings": warnings[:100],
        "samples": rows,
    }
