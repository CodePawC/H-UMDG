"""CSV / Excel helpers for imports."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_csv_dicts(data: bytes) -> list[dict[str, str]]:
    text = decode_text(data)
    reader = csv.DictReader(io.StringIO(text))
    rows: list[dict[str, str]] = []
    for raw in reader:
        row = {k.strip(): (v or "").strip() if k else "" for k, v in raw.items() if k}
        rows.append(row)
    return rows


def inspect_csv(data: bytes) -> dict[str, Any]:
    text = decode_text(data)
    reader = csv.reader(io.StringIO(text))
    try:
        headers = next(reader)
    except StopIteration:
        headers = []
        row_count = 0
    else:
        row_count = sum(1 for row in reader if any(str(value).strip() for value in row))
    return {
        "file_type": "CSV",
        "sheet_required": False,
        "default_sheet_name": None,
        "sheets": [],
        "row_count": row_count,
        "headers": [str(header).strip() for header in headers if str(header).strip()][:30],
    }


def inspect_excel(data: bytes, preferred: str | None = None) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        default_sheet_name = pick_sheet_name(wb, preferred)
        sheets: list[dict[str, Any]] = []
        for name in wb.sheetnames:
            ws = wb[name]
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = [(str(value).strip() if value is not None else "") for value in header_row]
            sheets.append(
                {
                    "name": name,
                    "row_count": max((ws.max_row or 1) - 1, 0),
                    "column_count": ws.max_column or 0,
                    "headers": [header for header in headers if header][:30],
                    "selected_by_default": name == default_sheet_name,
                }
            )
        return {
            "file_type": "EXCEL",
            "sheet_required": len(sheets) > 1,
            "default_sheet_name": default_sheet_name,
            "sheets": sheets,
            "row_count": next((sheet["row_count"] for sheet in sheets if sheet["name"] == default_sheet_name), 0),
            "headers": next((sheet["headers"] for sheet in sheets if sheet["name"] == default_sheet_name), []),
        }
    finally:
        wb.close()


def inspect_tabular_file(file_content: bytes, filename: str, sheet_name: str | None = None) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return inspect_csv(file_content)
    if suffix in {".xlsx", ".xlsm"}:
        return inspect_excel(file_content, sheet_name)
    return {
        "file_type": suffix.removeprefix(".").upper() or "UNKNOWN",
        "sheet_required": False,
        "default_sheet_name": None,
        "sheets": [],
        "row_count": None,
        "headers": [],
    }


def pick_sheet_name(workbook, preferred: str | None) -> str:
    names = workbook.sheetnames
    if preferred:
        if preferred in names:
            return preferred
        raise ValueError(f"sheet not found: {preferred}")
    if "Query1" in names:
        return "Query1"
    return names[0]


def read_excel_dicts(data: bytes, sheet_name: str | None = None) -> tuple[str, list[dict[str, Any]]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    name = pick_sheet_name(wb, sheet_name)
    ws = wb[name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        return name, []

    headers = [(str(c).strip() if c is not None else "") for c in header_row]
    out: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        item: dict[str, Any] = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            val = row[i] if i < len(row) else None
            if val is None:
                item[h] = ""
            elif isinstance(val, float) and val.is_integer():
                item[h] = str(int(val))
            else:
                item[h] = str(val).strip()
        out.append(item)
    wb.close()
    return name, out


def _row_to_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    item: dict[str, Any] = {}
    for i, h in enumerate(headers):
        if not h:
            continue
        val = row[i] if i < len(row) else None
        if val is None:
            item[h] = ""
        elif isinstance(val, float) and val.is_integer():
            item[h] = str(int(val))
        else:
            item[h] = str(val).strip()
    return item


def iter_csv_dicts_from_path(path: Path) -> tuple[str, int | None, Iterator[dict[str, Any]]]:
    def gen() -> Iterator[dict[str, Any]]:
        with path.open("rb") as fh:
            text = decode_text(fh.read())
        reader = csv.DictReader(io.StringIO(text))
        for raw in reader:
            yield {k.strip(): (v or "").strip() if k else "" for k, v in raw.items() if k}

    return "", None, gen()


def iter_excel_dicts_from_path(path: Path, sheet_name: str | None = None) -> tuple[str, int | None, Iterator[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    name = pick_sheet_name(wb, sheet_name)
    ws = wb[name]
    total_rows = max((ws.max_row or 1) - 1, 0)

    def gen() -> Iterator[dict[str, Any]]:
        try:
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                return
            headers = [(str(c).strip() if c is not None else "") for c in header_row]
            for row in rows_iter:
                if row is None or all(v is None or str(v).strip() == "" for v in row):
                    continue
                yield _row_to_dict(headers, row)
        finally:
            wb.close()

    return name, total_rows, gen()


def iter_tabular_rows_from_path(path: Path, sheet_name: str | None) -> tuple[str, int | None, Iterator[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix in {".csv"}:
        return iter_csv_dicts_from_path(path)
    if suffix in {".xlsx", ".xlsm"}:
        return iter_excel_dicts_from_path(path, sheet_name)
    raise ValueError(f"unsupported file type: {suffix}")


def load_tabular_rows(file_content: bytes, filename: str, sheet_name: str | None) -> tuple[str, list[dict[str, Any]]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".csv"}:
        return "", read_csv_dicts(file_content)
    if suffix in {".xlsx", ".xlsm"}:
        return read_excel_dicts(file_content, sheet_name)
    raise ValueError(f"unsupported file type: {suffix}")
