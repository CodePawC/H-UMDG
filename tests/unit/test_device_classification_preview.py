from pathlib import Path

from docx import Document

from app.services.import_pipeline import preview_device_classification_file
from app.models.tables import RefDeviceClassificationCatalog
from app.services.equipment_dictionary import is_invalid_device_row


def _write_device_adjustment_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("附件")
    doc.add_paragraph("《医疗器械分类目录》部分内容调整表")
    table = doc.add_table(rows=2, cols=15)
    top_headers = ["序号"] + ["《医疗器械分类目录》内容"] * 7 + ["调整后《医疗器械分类目录》内容"] * 7
    detail_headers = [
        "序号",
        "子 目录",
        "一级产品类别",
        "二级产品类别",
        "产品描述",
        "预期用途",
        "品名举例",
        "管理类别",
        "子 目录",
        "一级产品类别",
        "二级产品类别",
        "产品描述",
        "预期用途",
        "品名举例",
        "管理类别",
    ]
    for index, value in enumerate(top_headers):
        table.rows[0].cells[index].text = value
    for index, value in enumerate(detail_headers):
        table.rows[1].cells[index].text = value

    rows = [
        [
            "1",
            "02无源手术器械",
            "11手术器械-牵开器",
            "03扩张器",
            "无",
            "无",
            "无",
            "无",
            "02无源手术器械",
            "11手术器械-牵开器",
            "03扩张器",
            "通常由主操作鞘管、扩张器等组成。无菌提供。",
            "用于微创手术前建立操作空间。",
            "内窥镜通道扩张器",
            "II",
        ],
        [
            "2",
            "03神经和心血管手术器械",
            "05神经和心血管手术器械-夹",
            "02止血夹",
            "原产品描述",
            "原预期用途",
            "银夹",
            "Ⅲ",
            "03神经和心血管手术器械",
            "05神经和心血管手术器械-夹",
            "02止血夹",
            "调整后产品描述",
            "无变化",
            "无变化",
            "无变化",
        ],
        [
            "3",
            "无",
            "无",
            "无",
            "无",
            "无",
            "无",
            "无",
            "01有源手术器械",
            "01超声手术设备及附件",
            "04新增设备附件",
            "新增描述",
            "新增用途",
            "新增品名",
            "I",
        ],
    ]
    for value_row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(value_row):
            cells[index].text = value
    doc.save(path)


def test_preview_device_classification_adjustment_docx(tmp_path: Path) -> None:
    path = tmp_path / "国家药品监督管理局2022年第30号公告附件.docx"
    _write_device_adjustment_docx(path)

    result = preview_device_classification_file(path.read_bytes(), path.name)

    assert result["source_row_count"] == 3
    assert result["parse_success_count"] == 3
    assert result["parse_failed_count"] == 0
    assert result["batch_risk_level"] == "low"
    assert result["level_2_count"] == 3
    assert result["change_type_counts"]["新增"] == 1
    assert result["change_type_counts"]["修改"] == 2
    assert result["validation_results"][1]["status"] == "通过"
    assert result["samples"][0]["major_category_no"] == "02"
    assert result["samples"][0]["level_1_category_no"] == "11"
    assert result["samples"][0]["level_2_category"] == "扩张器"
    assert result["samples"][1]["intended_use"] == "原预期用途"
    assert result["samples"][1]["management_class"] == "Ⅲ"
    assert result["samples"][2]["change_type"] == "新增"


def test_preview_blocks_hyphen_prefixed_noise_row(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("01 有源手术器械")
    table = doc.add_table(rows=1, cols=7)
    headers = ["序号", "一级产品类别", "二级产品类别", "产品描述", "预期用途", "品名举例", "管理类别"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    row = table.add_row().cells
    values = ["01", "超声手术设备及附件", "-子目录", "产品描述", "预期用途", "品名举例", "II"]
    for index, value in enumerate(values):
        row[index].text = value
    path = tmp_path / "异常目录.docx"
    doc.save(path)

    result = preview_device_classification_file(path.read_bytes(), path.name)

    assert result["pending_confirm_count"] == 1
    assert result["invalid_count"] == 1
    assert result["batch_risk_level"] == "high"
    assert result["low_confidence_count"] == 1
    sample = result["samples"][0]
    assert sample["parse_status"] == "待人工确认"
    assert sample["data_status"] == "parse_abnormal"
    assert sample["abnormal_type"] == "解析噪声"
    assert sample["hidden_in_tree"] is True
    assert sample["validation_issue_codes"] == ["device_catalog_name_bad_prefix"]
    assert sample["confidence_score"] < 60


def test_governance_scan_does_not_flag_valid_adjustment_terms() -> None:
    row = RefDeviceClassificationCatalog(
        batch_id="TEST",
        source_file_name="国家药品监督管理局2022年第30号公告附件.docx",
        source_table_index=0,
        row_number=3,
        category_no="01",
        major_category_no="01",
        major_category_name="有源手术器械",
        level_1_category_no="01",
        level_1_category="超声手术设备及附件",
        level_2_category_no="04",
        level_2_category="新增设备附件",
        product_description="新增描述",
        intended_use="新增用途",
        product_examples="新增品名",
        management_class="I",
        data_status="effective",
        hidden_in_tree=False,
    )

    invalid, reasons = is_invalid_device_row(row)

    assert invalid is False
    assert reasons == []
