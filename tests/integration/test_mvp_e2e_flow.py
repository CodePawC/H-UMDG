from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import delete, func, or_, select, text
from sqlalchemy.exc import OperationalError

from app.db.session import get_session_factory
from app.main import app
from app.models.tables import (
    ApiCallLog,
    ApiClient,
    CatalogChangeHistory,
    CatalogCorrectionLog,
    CatalogCorrectionOrder,
    DictDepartment,
    DictEquipmentCategory,
    DictEquipmentStandardName,
    DictMaterialSpec,
    ManufacturerVendorCandidate,
    ManufacturerVendorExternalMapping,
    ManufacturerVendorMaster,
    ManufacturerVendorRelation,
    ManufacturerVendorRole,
    RefDeviceClassificationCatalog,
    RefDeviceClassificationRevision,
    StgNhsaMaterialDisabled,
    StgNhsaMaterialSpec,
    StgNhsaMaterialTranscode,
    SysExchangeLog,
    SysImportBatch,
    SysImportFailure,
    SysMappingBridge,
    SysMappingReviewTask,
)
from app.services.mapping_bridge import import_spd_mapping_csv
from app.services.standard_departments import STANDARD_DEPARTMENT_VERSION

REPO_ROOT = Path(__file__).resolve().parents[2]
HEADERS = {"X-API-Key": "change-me"}


def _db_available() -> bool:
    try:
        SessionLocal = get_session_factory()
        with SessionLocal() as db:
            db.execute(text("select 1"))
        return True
    except OperationalError:
        return False


requires_db = pytest.mark.skipif(not _db_available(), reason="PostgreSQL test database is not available")


def _clean_e2e_data() -> None:
    SessionLocal = get_session_factory()
    with SessionLocal() as db:
        db.execute(delete(SysExchangeLog).where(SysExchangeLog.source_tx_id.like("PY-E2E-%")))
        db.execute(delete(ApiCallLog).where(ApiCallLog.trace_id.like("req-py-e2e-%")))
        vendor_ids = select(ManufacturerVendorMaster.id).where(ManufacturerVendorMaster.organization_code.like("PY-E2E-%"))
        db.execute(
            delete(DictMaterialSpec).where(
                or_(
                    DictMaterialSpec.yb_code_27.in_(["123456789012345678901234567", "123456789012345678901234568"]),
                    DictMaterialSpec.source_batch_id.like("PY-E2E-%"),
                    DictMaterialSpec.manufacturer_org_id.in_(vendor_ids),
                    DictMaterialSpec.registrant_org_id.in_(vendor_ids),
                    DictMaterialSpec.filer_org_id.in_(vendor_ids),
                )
            )
        )
        db.execute(
            delete(ManufacturerVendorExternalMapping).where(
                or_(ManufacturerVendorExternalMapping.external_code.like("PY-E2E-%"), ManufacturerVendorExternalMapping.org_id.in_(vendor_ids))
            )
        )
        db.execute(delete(ManufacturerVendorCandidate).where(ManufacturerVendorCandidate.source_record_id.like("PY-E2E-%")))
        db.execute(
            delete(ManufacturerVendorRelation).where(
                or_(
                    ManufacturerVendorRelation.remark.like("PY-E2E-%"),
                    ManufacturerVendorRelation.parent_org_id.in_(vendor_ids),
                    ManufacturerVendorRelation.child_org_id.in_(vendor_ids),
                )
            )
        )
        db.execute(
            delete(ManufacturerVendorRole).where(
                or_(ManufacturerVendorRole.remark.like("PY-E2E-%"), ManufacturerVendorRole.org_id.in_(vendor_ids))
            )
        )
        db.execute(delete(ManufacturerVendorMaster).where(ManufacturerVendorMaster.organization_code.like("PY-E2E-%")))
        db.execute(delete(SysImportFailure).where(SysImportFailure.batch_id.like("PY-E2E-%")))
        device_catalog_ids = select(RefDeviceClassificationCatalog.catalog_id).where(
            RefDeviceClassificationCatalog.batch_id.like("PY-E2E-%")
        )
        correction_ids = select(CatalogCorrectionOrder.correction_id).where(
            or_(
                CatalogCorrectionOrder.source_batch_id.like("PY-E2E-%"),
                CatalogCorrectionOrder.catalog_id.in_(device_catalog_ids),
            )
        )
        db.execute(delete(CatalogCorrectionLog).where(CatalogCorrectionLog.correction_id.in_(correction_ids)))
        db.execute(delete(CatalogChangeHistory).where(CatalogChangeHistory.correction_id.in_(correction_ids)))
        db.execute(
            delete(CatalogCorrectionOrder).where(
                or_(
                    CatalogCorrectionOrder.source_batch_id.like("PY-E2E-%"),
                    CatalogCorrectionOrder.catalog_id.in_(device_catalog_ids),
                )
            )
        )
        db.execute(delete(SysImportBatch).where(SysImportBatch.batch_id.like("PY-E2E-%")))
        db.execute(delete(SysMappingBridge).where(SysMappingBridge.source_system == "SPD"))
        db.execute(delete(SysMappingReviewTask).where(SysMappingReviewTask.source_system == "SPD"))
        db.execute(delete(StgNhsaMaterialTranscode).where(StgNhsaMaterialTranscode.batch_id.like("PY-E2E-%")))
        db.execute(delete(StgNhsaMaterialDisabled).where(StgNhsaMaterialDisabled.batch_id.like("PY-E2E-%")))
        db.execute(delete(StgNhsaMaterialSpec).where(StgNhsaMaterialSpec.batch_id.like("PY-E2E-%")))
        db.execute(
            delete(RefDeviceClassificationRevision).where(RefDeviceClassificationRevision.batch_id.like("PY-E2E-%"))
        )
        db.execute(
            delete(RefDeviceClassificationCatalog).where(RefDeviceClassificationCatalog.batch_id.like("PY-E2E-%"))
        )
        db.execute(delete(DictEquipmentStandardName).where(DictEquipmentStandardName.source_batch_id.like("PY-E2E-%")))
        db.execute(delete(DictEquipmentCategory).where(DictEquipmentCategory.source_batch_id.like("PY-E2E-%")))
        db.execute(delete(DictDepartment).where(DictDepartment.dept_code.in_(["DEPT001", "DEPT002"])))
        db.execute(delete(DictDepartment).where(DictDepartment.source_batch_id == STANDARD_DEPARTMENT_VERSION))
        db.commit()


def _post_file(client: TestClient, url: str, path: Path, fields: dict[str, str]):
    with path.open("rb") as fh:
        return client.post(url, headers=HEADERS, data=fields, files={"file": (path.name, fh, "text/csv")})


def _write_device_catalog_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("01 有源手术器械说明")
    doc.add_paragraph("说明段落")
    doc.add_paragraph("01 有源手术器械")
    ignored = doc.add_table(rows=2, cols=2)
    ignored.rows[0].cells[0].text = "not"
    ignored.rows[0].cells[1].text = "catalog"
    ignored.rows[1].cells[0].text = "ignored"
    ignored.rows[1].cells[1].text = "ignored"

    table = doc.add_table(rows=1, cols=7)
    headers = ["序号", "一级产品类别", "二级产品类别", "产品描述", "预期用途", "品名举例", "管理类别"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    values = [
        ["01", "超声手术设备及附件", "01超声手术设备", "用于手术操作", "临床手术", "高频手术设备", "II"],
        ["02", "激光手术设备及附件", "01激光手术设备", "用于骨修复", "骨科治疗", "接骨板", "III"],
        ["03", "手术器械附件", "03 扩孔 用刀", "用于扩孔 操作", "骨科 手术", "扩孔 用刀", "II"],
        ["", "", "", "", "", "", ""],
    ]
    for value_row in values:
        cells = table.add_row().cells
        for i, value in enumerate(value_row):
            cells[i].text = value
    doc.save(path)


def _write_device_adjustment_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("《医疗器械分类目录》部分内容调整表")
    table = doc.add_table(rows=2, cols=15)
    top_headers = ["序号"] + ["《医疗器械分类目录》内容"] * 7 + ["调整后《医疗器械分类目录》内容"] * 7
    detail_headers = [
        "序号",
        "子目录",
        "一级产品类别",
        "二级产品类别",
        "产品描述",
        "预期用途",
        "品名举例",
        "管理类别",
        "子 目 录",
        "一级产品类别",
        "二级产品类别",
        "产品描述",
        "预期用途",
        "品名举例",
        "管理类别",
    ]
    for index, value in enumerate(top_headers):
        table.rows[0].cells[index].text = value
    for index, value in enumerate(detail_headers):
        table.rows[1].cells[index].text = value
    values = [
        [
            "1",
            "01有源手术器械",
            "01超声手术设备及附件",
            "01超声手术设备",
            "用于手术操作",
            "临床手术",
            "高频手术设备",
            "II",
            "01有源手术器械",
            "01超声手术设备及附件",
            "01超声手术设备",
            "用于手术操作调整",
            "无变化",
            "无变化",
            "无变化",
        ],
        [
            "2",
            "无",
            "无",
            "无",
            "无",
            "无",
            "无",
            "无",
            "01有源手术器械",
            "01超声手术设备及附件",
            "04新增设备附件",
            "新增描述",
            "新增用途",
            "新增品名",
            "I",
        ],
    ]
    for value_row in values:
        cells = table.add_row().cells
        for index, value in enumerate(value_row):
            cells[index].text = value
    doc.save(path)


@requires_db
def test_mvp_department_material_spd_mapping_flow() -> None:
    _clean_e2e_data()
    client = TestClient(app)

    dept_path = REPO_ROOT / "data/templates/H-UDMP-TEMPLATE-DEPARTMENTS-v1.0.csv"
    mat_path = REPO_ROOT / "data/templates/H-UDMP-TEMPLATE-MATERIALS-v1.0.csv"
    spd_path = REPO_ROOT / "data/samples/H-UDMP-SAMPLE-SPD-MAPPINGS-v1.0.csv"

    dept_resp = _post_file(
        client,
        "/api/v1/departments/import",
        dept_path,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-DEPT-001"},
    )
    assert dept_resp.status_code == 200
    assert dept_resp.json()["data"]["success_count"] == 2
    assert dept_resp.json()["data"]["failed_count"] == 0

    mat_resp = _post_file(
        client,
        "/api/v1/materials/import",
        mat_path,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-MAT-001", "source_type": "MVP_TEMPLATE"},
    )
    assert mat_resp.status_code == 200
    assert mat_resp.json()["data"]["success_count"] == 2
    assert mat_resp.json()["data"]["failed_count"] == 0

    SessionLocal = get_session_factory()
    with SessionLocal() as db:
        report = import_spd_mapping_csv(db, spd_path.read_bytes(), spd_path.name, "PY-E2E-SPD-001")
        db.commit()
    assert report["success_count"] == 2
    assert report["failed_count"] == 0

    mat_resolve = client.post(
        "/api/v1/mapping/resolve",
        headers=HEADERS,
        json={
            "category": "material",
            "source_system": "SPD",
            "source_key": "SPD-MAT-001",
            "source_desc": "一次性使用无菌导管 22mm",
            "source_tx_id": "PY-E2E-MAP-MAT-001",
        },
    )
    assert mat_resolve.status_code == 200
    assert mat_resolve.json()["data"]["status"] == "MATCHED"
    assert mat_resolve.json()["data"]["target_code"] == "123456789012345678901234567"

    dept_resolve = client.post(
        "/api/v1/mapping/resolve",
        headers=HEADERS,
        json={
            "category": "department",
            "source_system": "SPD",
            "source_key": "SPD-DEPT-001",
            "source_desc": "心外科",
            "source_tx_id": "PY-E2E-MAP-DEPT-001",
        },
    )
    assert dept_resolve.status_code == 200
    assert dept_resolve.json()["data"]["status"] == "MATCHED"
    assert dept_resolve.json()["data"]["target_code"] == "DEPT001"

    alias_search = client.get("/api/v1/departments/search?keyword=心外", headers=HEADERS)
    assert alias_search.status_code == 200
    assert alias_search.json()["data"]["page"]["total"] == 1
    dept_import_preview = client.get(
        "/api/v1/departments/search?source_batch_id=PY-E2E-DEPT-001&status=",
        headers=HEADERS,
    )
    assert dept_import_preview.status_code == 200
    assert dept_import_preview.json()["data"]["page"]["total"] == 2

    material_search = client.get(
        "/api/v1/materials/search?yb_code_27=123456789012345678901234567",
        headers=HEADERS,
    )
    assert material_search.status_code == 200
    assert material_search.json()["data"]["page"]["total"] == 1
    material_import_preview = client.get(
        "/api/v1/materials/search?source_batch_id=PY-E2E-MAT-001&status=",
        headers=HEADERS,
    )
    assert material_import_preview.status_code == 200
    assert material_import_preview.json()["data"]["page"]["total"] == 2

    log_resp = client.get("/api/v1/exchange/logs?source_system=SPD&page_size=20", headers=HEADERS)
    assert log_resp.status_code == 200
    assert log_resp.json()["data"]["page"]["total"] >= 2
    mapping_log_resp = client.get("/api/v1/exchange/logs?source_system=SPD&data_category=mapping&page_size=20", headers=HEADERS)
    assert mapping_log_resp.status_code == 200
    mapping_logs = mapping_log_resp.json()["data"]["items"]
    assert mapping_logs
    assert all(item["data_category"] == "mapping" for item in mapping_logs)

    detail_resp = client.get(f"/api/v1/exchange/logs/{mapping_logs[0]['log_id']}", headers=HEADERS)
    assert detail_resp.status_code == 200
    detail = detail_resp.json()["data"]
    assert detail["payload_policy"]["raw_request_payload_included"] is False
    assert detail["payload_policy"]["raw_response_payload_included"] is False
    assert "request_payload" not in detail
    assert "response_payload" not in detail
    assert "request_payload_summary" in detail
    assert "response_payload_summary" in detail


@requires_db
def test_standard_department_library_sync_and_status_toggle() -> None:
    _clean_e2e_data()
    client = TestClient(app)

    sync_resp = client.post("/api/v1/departments/standard-library/sync", headers=HEADERS)
    assert sync_resp.status_code == 200
    sync_data = sync_resp.json()["data"]
    assert sync_data["source_type"] == "STANDARD_DEPARTMENT_LIBRARY"
    assert sync_data["success_count"] >= 100
    assert sync_data["failed_count"] == 0

    search_resp = client.get("/api/v1/departments/search?keyword=心血管&status=", headers=HEADERS)
    assert search_resp.status_code == 200
    rows = search_resp.json()["data"]["items"]
    target = next(row for row in rows if row["dept_code"] == "NHC-03.04")
    assert target["dept_name"] == "心血管内科专业"
    assert target["standard_scope"]
    assert target["maintenance_mode"] == "STANDARD_LIBRARY_TOGGLE"

    toggle_resp = client.patch(
        "/api/v1/departments/NHC-03.04/status",
        headers=HEADERS,
        json={"status": "INACTIVE"},
    )
    assert toggle_resp.status_code == 200
    assert toggle_resp.json()["data"]["status"] == "INACTIVE"

    inactive_resp = client.get("/api/v1/departments/search?keyword=心血管&status=INACTIVE", headers=HEADERS)
    assert inactive_resp.status_code == 200
    assert inactive_resp.json()["data"]["page"]["total"] >= 1

    restore_resp = client.patch(
        "/api/v1/departments/NHC-03.04/status",
        headers=HEADERS,
        json={"status": "ACTIVE"},
    )
    assert restore_resp.status_code == 200


@requires_db
def test_department_import_idempotency_replay_and_conflict(tmp_path: Path) -> None:
    _clean_e2e_data()
    client = TestClient(app)
    dept_path = REPO_ROOT / "data/templates/H-UDMP-TEMPLATE-DEPARTMENTS-v1.0.csv"

    first = _post_file(
        client,
        "/api/v1/departments/import",
        dept_path,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-IDEMPOTENCY-001"},
    )
    replay = _post_file(
        client,
        "/api/v1/departments/import",
        dept_path,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-IDEMPOTENCY-001"},
    )
    assert first.status_code == 200
    assert replay.status_code == 200
    assert replay.json()["data"] == first.json()["data"]

    changed = tmp_path / "departments-changed.csv"
    changed.write_text(
        dept_path.read_text(encoding="utf-8").replace("心脏外科", "心脏外科门诊"),
        encoding="utf-8",
    )
    conflict = _post_file(
        client,
        "/api/v1/departments/import",
        changed,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-IDEMPOTENCY-001"},
    )
    assert conflict.status_code == 409
    assert conflict.json()["code"] == "IDEMPOTENCY_CONFLICT"


@requires_db
def test_mapping_review_approve_reject_flow() -> None:
    _clean_e2e_data()
    client = TestClient(app)
    mat_path = REPO_ROOT / "data/templates/H-UDMP-TEMPLATE-MATERIALS-v1.0.csv"

    mat_resp = _post_file(
        client,
        "/api/v1/materials/import",
        mat_path,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-REVIEW-MAT-001", "source_type": "MVP_TEMPLATE"},
    )
    assert mat_resp.status_code == 200

    pending_resp = client.post(
        "/api/v1/mapping/resolve",
        headers=HEADERS,
        json={
            "category": "material",
            "source_system": "SPD",
            "source_key": "SPD-MAT-REVIEW-001",
            "source_desc": "待审核耗材",
            "source_tx_id": "PY-E2E-REVIEW-RESOLVE-001",
        },
    )
    assert pending_resp.status_code == 200
    pending_data = pending_resp.json()["data"]
    assert pending_data["status"] == "PENDING_REVIEW"
    task_id = pending_data["task_id"]

    list_pending = client.get("/api/v1/mapping/review-tasks?source_system=SPD&status=PENDING", headers=HEADERS)
    assert list_pending.status_code == 200
    assert any(item["task_id"] == task_id for item in list_pending.json()["data"]["items"])

    SessionLocal = get_session_factory()
    with SessionLocal() as db:
        material = db.scalar(
            text("select material_id from dict_material_specs where yb_code_27='123456789012345678901234567'")
        )
    assert material is not None

    approve_resp = client.post(
        f"/api/v1/mapping/review-tasks/{task_id}/approve",
        headers=HEADERS,
        json={
            "target_master_id": str(material),
            "reviewer": "data_admin",
            "review_comment": "确认命中标准耗材",
            "mapping_rule": "MANUAL",
            "confidence": 0.95,
        },
    )
    assert approve_resp.status_code == 200
    assert approve_resp.json()["data"]["status"] == "APPROVED"

    matched_resp = client.post(
        "/api/v1/mapping/resolve",
        headers=HEADERS,
        json={
            "category": "material",
            "source_system": "SPD",
            "source_key": "SPD-MAT-REVIEW-001",
            "source_desc": "待审核耗材",
            "source_tx_id": "PY-E2E-REVIEW-RESOLVE-002",
        },
    )
    assert matched_resp.status_code == 200
    assert matched_resp.json()["data"]["status"] == "MATCHED"
    assert matched_resp.json()["data"]["target_code"] == "123456789012345678901234567"

    repeat_approve = client.post(
        f"/api/v1/mapping/review-tasks/{task_id}/approve",
        headers=HEADERS,
        json={"target_master_id": str(material), "reviewer": "data_admin"},
    )
    assert repeat_approve.status_code == 409
    assert repeat_approve.json()["code"] == "TASK_ALREADY_REVIEWED"

    reject_seed = client.post(
        "/api/v1/mapping/resolve",
        headers=HEADERS,
        json={
            "category": "material",
            "source_system": "SPD",
            "source_key": "SPD-MAT-REJECT-001",
            "source_desc": "无效耗材编码",
            "source_tx_id": "PY-E2E-REVIEW-REJECT-SEED-001",
        },
    )
    assert reject_seed.status_code == 200
    reject_task_id = reject_seed.json()["data"]["task_id"]

    reject_resp = client.post(
        f"/api/v1/mapping/review-tasks/{reject_task_id}/reject",
        headers=HEADERS,
        json={"reviewer": "data_admin", "review_comment": "来源编码无效"},
    )
    assert reject_resp.status_code == 200
    assert reject_resp.json()["data"]["status"] == "REJECTED"
    assert reject_resp.json()["data"]["reviewer"] == "data_admin"


@requires_db
def test_nhsa_sample_workbooks_import() -> None:
    _clean_e2e_data()
    client = TestClient(app)
    nhsa_dir = REPO_ROOT / "data/samples/external/nhsa"

    full_spec = nhsa_dir / "H-UDMP-SAMPLE-NHSA-C09-FULL-SPEC-v1.0.xlsx"
    disabled = nhsa_dir / "H-UDMP-SAMPLE-NHSA-DISABLED-v1.0.xlsx"
    transcode = nhsa_dir / "H-UDMP-SAMPLE-NHSA-TRANSCODE-v1.0.xlsx"

    full_resp = _post_file(
        client,
        "/api/v1/materials/import",
        full_spec,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-NHSA-FULL-001",
            "source_type": "NHSA_FULL_SPEC",
            "sheet_name": "Query1",
        },
    )
    assert full_resp.status_code == 200
    assert full_resp.json()["data"]["sheet_name"] == "Query1"
    assert full_resp.json()["data"]["source_row_count"] == 12
    assert full_resp.json()["data"]["unique_key_count"] == 12
    assert full_resp.json()["data"]["skipped_duplicate_count"] == 0
    assert full_resp.json()["data"]["success_count"] == 12
    assert full_resp.json()["data"]["failed_count"] == 0

    disabled_resp = _post_file(
        client,
        "/api/v1/materials/import",
        disabled,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-NHSA-DISABLED-001",
            "source_type": "NHSA_DISABLED",
            "sheet_name": "Query1",
        },
    )
    assert disabled_resp.status_code == 200
    assert disabled_resp.json()["data"]["source_row_count"] == 12
    assert disabled_resp.json()["data"]["unique_key_count"] == 12
    assert disabled_resp.json()["data"]["skipped_duplicate_count"] == 0
    assert disabled_resp.json()["data"]["success_count"] == 12
    assert disabled_resp.json()["data"]["disabled_count"] == 12
    assert disabled_resp.json()["data"]["failed_count"] == 0

    transcode_resp = _post_file(
        client,
        "/api/v1/materials/import",
        transcode,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-NHSA-TRANSCODE-001",
            "source_type": "NHSA_TRANSCODE",
            "sheet_name": "Query1",
        },
    )
    assert transcode_resp.status_code == 200
    assert transcode_resp.json()["data"]["source_row_count"] == 12
    assert transcode_resp.json()["data"]["unique_key_count"] == 12
    assert transcode_resp.json()["data"]["skipped_duplicate_count"] == 0
    assert transcode_resp.json()["data"]["success_count"] == 12
    assert transcode_resp.json()["data"]["transcoded_count"] == 12
    assert transcode_resp.json()["data"]["changed_count"] + transcode_resp.json()["data"]["mapped_count"] == 12
    assert transcode_resp.json()["data"]["failed_count"] == 0

    SessionLocal = get_session_factory()
    with SessionLocal() as db:
        assert db.scalar(text("select count(*) from stg_nhsa_material_specs where batch_id='PY-E2E-NHSA-FULL-001'")) == 12
        assert db.scalar(text("select count(*) from stg_nhsa_material_disabled where batch_id='PY-E2E-NHSA-DISABLED-001'")) == 12
        assert db.scalar(text("select count(*) from stg_nhsa_material_transcode where batch_id='PY-E2E-NHSA-TRANSCODE-001'")) == 12
        assert db.scalar(text("select count(*) from dict_material_specs where source_batch_id='PY-E2E-NHSA-FULL-001'")) == 12
        transcode_row = db.execute(
            text(
                "select original_yb_code_27, yb_code_27, change_type "
                "from stg_nhsa_material_transcode "
                "where batch_id='PY-E2E-NHSA-TRANSCODE-001' "
                "order by row_number limit 1"
            )
        ).mappings().one()

    resolve_resp = client.get(
        "/api/v1/materials/transcode/resolve",
        headers=HEADERS,
        params={
            "original_yb_code_27": transcode_row["original_yb_code_27"],
            "batch_id": "PY-E2E-NHSA-TRANSCODE-001",
        },
    )
    assert resolve_resp.status_code == 200
    resolve_data = resolve_resp.json()["data"]
    assert resolve_data["status"] == "FOUND"
    assert resolve_data["page"]["total"] == 1
    assert resolve_data["items"][0]["yb_code_27"] == transcode_row["yb_code_27"]
    assert resolve_data["items"][0]["change_type"] == transcode_row["change_type"]

    preview_resp = client.get(
        "/api/v1/materials/search?source_batch_id=PY-E2E-NHSA-FULL-001&status=",
        headers=HEADERS,
    )
    assert preview_resp.status_code == 200
    assert preview_resp.json()["data"]["page"]["total"] == 12

    not_found_resp = client.get(
        "/api/v1/materials/transcode/resolve",
        headers=HEADERS,
        params={"original_yb_code_27": "Z99999999999999999999999999"},
    )
    assert not_found_resp.status_code == 200
    assert not_found_resp.json()["data"]["status"] == "NOT_FOUND"
    assert not_found_resp.json()["data"]["page"]["total"] == 0


@requires_db
def test_device_classification_catalog_import(tmp_path: Path) -> None:
    _clean_e2e_data()
    client = TestClient(app)
    catalog_path = tmp_path / "device-catalog.docx"
    _write_device_catalog_docx(catalog_path)

    resp = _post_file(
        client,
        "/api/v1/materials/import",
        catalog_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-DEVICE-CATALOG-001",
            "source_type": "DEVICE_CLASSIFICATION_CATALOG",
        },
    )
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert data["source_type"] == "DEVICE_CLASSIFICATION_CATALOG"
    assert data["source_row_count"] == 3
    assert data["unique_key_count"] == 3
    assert data["success_count"] == 3
    assert data["failed_count"] == 0

    missing_detail_path = tmp_path / "device-catalog-missing-detail.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "序号"
    table.rows[0].cells[1].text = "一级产品类别"
    table.rows[1].cells[0].text = "01"
    table.rows[1].cells[1].text = "不完整目录"
    doc.save(missing_detail_path)
    missing_resp = _post_file(
        client,
        "/api/v1/materials/import",
        missing_detail_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-DEVICE-CATALOG-MISSING-001",
            "source_type": "DEVICE_CLASSIFICATION_CATALOG",
        },
    )
    assert missing_resp.status_code == 200
    missing_data = missing_resp.json()["data"]
    assert missing_data["success_count"] == 0
    assert missing_data["failed_count"] == 1
    assert missing_data["failures"][0]["field"] == "docx"

    SessionLocal = get_session_factory()
    with SessionLocal() as db:
        rows = db.execute(
            text(
                "select category_no, major_category_no, major_category_name, "
                "level_1_category_no, level_1_category, level_2_category_no, level_2_category, "
                "product_description, intended_use, product_examples, management_class "
                "from ref_device_classification_catalog "
                "where batch_id='PY-E2E-DEVICE-CATALOG-001' "
                "order by row_number"
            )
        ).mappings().all()
    assert [row["category_no"] for row in rows] == ["01", "02", "03"]
    assert rows[0]["major_category_no"] == "01"
    assert rows[0]["major_category_name"] == "有源手术器械"
    assert rows[0]["level_1_category_no"] == "01"
    assert rows[0]["level_1_category"] == "超声手术设备及附件"
    assert rows[0]["level_2_category_no"] == "01"
    assert rows[0]["level_2_category"] == "超声手术设备"
    assert rows[1]["level_1_category_no"] == "02"
    assert rows[1]["management_class"] == "III"
    assert rows[2]["level_2_category_no"] == "03"
    assert rows[2]["level_2_category"] == "扩孔用刀"
    assert rows[2]["product_description"] == "用于扩孔操作"
    assert rows[2]["intended_use"] == "骨科手术"
    assert rows[2]["product_examples"] == "扩孔用刀"

    adjustment_path = tmp_path / "device-catalog-adjustment.docx"
    _write_device_adjustment_docx(adjustment_path)
    adjustment_resp = _post_file(
        client,
        "/api/v1/materials/import",
        adjustment_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-DEVICE-CATALOG-ADJUST-001",
            "source_type": "DEVICE_CLASSIFICATION_CATALOG",
        },
    )
    assert adjustment_resp.status_code == 200
    adjustment_data = adjustment_resp.json()["data"]
    assert adjustment_data["source_row_count"] == 2
    assert adjustment_data["success_count"] == 2
    assert adjustment_data["changed_count"] == 1

    with SessionLocal() as db:
        current_rows = db.execute(
            text(
                "select catalog_id, level_2_category_no, level_2_category, product_description, intended_use, "
                "product_examples, management_class, data_status, hidden_in_tree from ref_device_classification_catalog "
                "where major_category_no='01' and level_1_category_no='01' "
                "order by level_2_category_no, product_description"
            )
        ).mappings().all()
    old_rows = [row for row in current_rows if row["product_description"] == "用于手术操作"]
    assert old_rows
    assert all(row["data_status"] == "deprecated" and row["hidden_in_tree"] for row in old_rows)
    descriptions = {row["product_description"] for row in current_rows if row["data_status"] == "effective"}
    assert "用于手术操作调整" in descriptions
    added = next(row for row in current_rows if row["level_2_category_no"] == "04")
    assert added["level_2_category"] == "新增设备附件"
    assert added["product_examples"] == "新增品名"

    tree_resp = client.get("/api/v1/equipment/device-classifications/tree-data?keyword=新增设备附件", headers=HEADERS)
    assert tree_resp.status_code == 200
    tree_data = tree_resp.json()["data"]
    tree_items = tree_data["items"]
    assert tree_items[0]["latest_revision"]["change_type"] == "ADDED"
    assert tree_data["revision_summary"]["revised_item_count"] >= 2
    assert tree_data["revision_summary"]["last_updated_at"]

    revisions_resp = client.get(
        f"/api/v1/equipment/device-classifications/revisions?catalog_id={added['catalog_id']}",
        headers=HEADERS,
    )
    assert revisions_resp.status_code == 200
    revisions = revisions_resp.json()["data"]["items"]
    assert revisions[0]["change_type"] == "ADDED"
    assert revisions[0]["new_payload"]["product_examples"] == "新增品名"


@requires_db
def test_device_classification_source_trace_download_and_correction(tmp_path: Path) -> None:
    _clean_e2e_data()
    client = TestClient(app)
    catalog_path = tmp_path / "device-catalog-source-trace.docx"
    _write_device_catalog_docx(catalog_path)

    resp = _post_file(
        client,
        "/api/v1/equipment/import",
        catalog_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-DEVICE-SOURCE-TRACE-001",
            "source_type": "DEVICE_CLASSIFICATION_CATALOG",
            "operator_name": "数据治理员",
            "import_mode": "增量更新",
        },
    )
    assert resp.status_code == 200

    SessionLocal = get_session_factory()
    with SessionLocal() as db:
        catalog_id = db.scalar(
            select(RefDeviceClassificationCatalog.catalog_id).where(
                RefDeviceClassificationCatalog.batch_id == "PY-E2E-DEVICE-SOURCE-TRACE-001"
            )
        )
    assert catalog_id is not None

    governance = client.get(
        f"/api/v1/equipment/device-classifications/{catalog_id}/governance",
        headers=HEADERS,
    )
    assert governance.status_code == 200
    trace = governance.json()["data"]["source_trace"]
    assert trace["source_batch_id"] == "PY-E2E-DEVICE-SOURCE-TRACE-001"
    assert trace["source_file_name"] == catalog_path.name
    assert trace["source_file_hash"]
    assert "第" in trace["source_position"]

    download = client.get(
        f"/api/v1/equipment/device-classifications/{catalog_id}/source-file/download",
        headers=HEADERS,
    )
    assert download.status_code == 200
    assert download.content == catalog_path.read_bytes()

    correction = client.post(
        f"/api/v1/equipment/device-classifications/{catalog_id}/corrections",
        headers=HEADERS,
        json={
            "abnormal_type": "目录名称异常",
            "correction_action": "保持不变",
            "reason": "PY-E2E 回归测试确认纠错闭环可提交",
            "hide_in_tree": False,
            "need_review": True,
        },
    )
    assert correction.status_code == 200
    correction_data = correction.json()["data"]
    assert correction_data["status"] == "submitted"
    assert correction_data["source_batch_id"] == "PY-E2E-DEVICE-SOURCE-TRACE-001"

    approved = client.post(
        f"/api/v1/equipment/device-classifications/corrections/{correction_data['correction_id']}/approve",
        headers=HEADERS,
    )
    assert approved.status_code == 200
    assert approved.json()["data"]["status"] == "approved"

    ledger = client.get(
        "/api/v1/equipment/device-classifications/corrections?status=approved&source_batch_id=PY-E2E-DEVICE-SOURCE-TRACE-001",
        headers=HEADERS,
    )
    assert ledger.status_code == 200
    assert ledger.json()["data"]["total"] == 1

    templates = client.get("/api/v1/equipment/device-classifications/export/templates", headers=HEADERS)
    assert templates.status_code == 200
    assert any(item["template_id"] == "governance" for item in templates.json()["data"]["items"])

    estimate = client.post(
        "/api/v1/equipment/device-classifications/export/estimate",
        headers=HEADERS,
        json={
            "scope_type": "batch",
            "template_id": "governance",
            "criteria": {"scope_type": "batch", "batch_id": "PY-E2E-DEVICE-SOURCE-TRACE-001"},
        },
    )
    assert estimate.status_code == 200
    assert estimate.json()["data"]["estimated_count"] >= 1

    export_resp = client.post(
        "/api/v1/equipment/device-classifications/export/tasks",
        headers=HEADERS,
        json={
            "task_name": "PY-E2E 医疗器械分类目录导出",
            "purpose": "回归测试",
            "scope_type": "batch",
            "criteria": {"scope_type": "batch", "batch_id": "PY-E2E-DEVICE-SOURCE-TRACE-001"},
            "template_id": "governance",
            "format": "xlsx",
        },
    )
    assert export_resp.status_code == 200
    export_data = export_resp.json()["data"]
    assert export_data["task_no"].startswith("EXP")
    assert export_data["row_count"] >= 1
    assert export_data["file_name"].endswith(".xlsx")
    assert export_data["file_hash"]

    export_download = client.get(
        f"/api/v1/equipment/device-classifications/export/tasks/{export_data['task_no']}/download",
        headers=HEADERS,
    )
    assert export_download.status_code == 200
    assert export_download.content[:2] == b"PK"

    export_detail = client.get(
        f"/api/v1/equipment/device-classifications/export/tasks/{export_data['task_no']}",
        headers=HEADERS,
    )
    assert export_detail.status_code == 200
    assert export_detail.json()["data"]["download_count"] >= 1


@requires_db
def test_master_data_device_classification_external_api_contract(tmp_path: Path) -> None:
    _clean_e2e_data()
    client = TestClient(app)
    catalog_path = tmp_path / "device-catalog-external-api.docx"
    _write_device_catalog_docx(catalog_path)
    import_resp = _post_file(
        client,
        "/api/v1/equipment/import",
        catalog_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-MD-API-001",
            "source_type": "DEVICE_CLASSIFICATION_CATALOG",
            "operator_name": "数据治理员",
        },
    )
    assert import_resp.status_code == 200

    SessionLocal = get_session_factory()
    with SessionLocal() as db:
        db.add(
            RefDeviceClassificationCatalog(
                batch_id="PY-E2E-MD-API-001",
                source_file_name="abnormal.docx",
                source_table_index=1,
                row_number=99,
                category_no="PY-ABNORMAL",
                major_category_no="99",
                major_category_name="-子目录",
                level_1_category_no="99",
                level_1_category="-子目录",
                level_2_category_no="99",
                level_2_category="-子目录",
                management_class="II",
                data_status="parse_abnormal",
                status_reason="PY-E2E 异常树过滤测试",
                hidden_in_tree=False,
            )
        )
        db.commit()
        effective = db.scalar(
            select(RefDeviceClassificationCatalog).where(
                RefDeviceClassificationCatalog.batch_id == "PY-E2E-MD-API-001",
                RefDeviceClassificationCatalog.data_status == "effective",
            )
        )
        assert effective is not None
        effective_id = str(effective.catalog_id)
        effective_code = effective.category_no or effective.level_2_category_no or effective.level_1_category_no or effective.major_category_no

    api_headers = {
        "Authorization": "Bearer hudmp-dbus-dev-20260525",
        "User-Agent": "PY-E2E-MD-CLIENT",
    }

    missing_key = client.get("/api/v1/master-data/device-classification/catalog?keyword=手术")
    assert missing_key.status_code == 401
    assert missing_key.json()["success"] is False
    assert missing_key.json()["code"] == "API_KEY_MISSING"

    wrong_key = client.get(
        "/api/v1/master-data/device-classification/catalog?keyword=手术",
        headers={"X-API-Key": "wrong-key", "X-Request-ID": "req-py-e2e-md-wrong"},
    )
    assert wrong_key.status_code == 401
    assert wrong_key.json()["code"] == "API_KEY_INVALID"

    catalog = client.get(
        "/api/v1/master-data/device-classification/catalog?keyword=手术&page=1&pageSize=10",
        headers=api_headers | {"X-Request-ID": "req-py-e2e-md-catalog"},
    )
    assert catalog.status_code == 200
    catalog_payload = catalog.json()
    assert catalog_payload["success"] is True
    assert catalog_payload["traceId"] == "req-py-e2e-md-catalog"
    records = catalog_payload["data"]["records"]
    assert records
    assert {row["status"] for row in records} == {"effective"}
    assert all(row["name"] != "-子目录" for row in records)

    detail = client.get(
        f"/api/v1/master-data/device-classification/catalog/{effective_id}",
        headers=api_headers | {"X-Request-ID": "req-py-e2e-md-detail"},
    )
    assert detail.status_code == 200
    assert detail.json()["data"]["id"] == effective_id

    by_code = client.get(
        f"/api/v1/master-data/device-classification/catalog/by-code/{effective_code}",
        headers=api_headers | {"X-Request-ID": "req-py-e2e-md-code"},
    )
    assert by_code.status_code == 200
    assert by_code.json()["data"]["code"] == effective_code

    tree = client.get(
        "/api/v1/master-data/device-classification/tree?keyword=手术&maxDepth=3",
        headers=api_headers | {"X-Request-ID": "req-py-e2e-md-tree"},
    )
    assert tree.status_code == 200
    assert "-子目录" not in str(tree.json()["data"]["records"])

    current_version = client.get(
        "/api/v1/master-data/device-classification/current-version",
        headers=api_headers | {"X-Request-ID": "req-py-e2e-md-version"},
    )
    assert current_version.status_code == 200
    assert current_version.json()["data"]["status"] == "active"

    changes = client.get(
        "/api/v1/master-data/device-classification/changes?since=2000-01-01T00:00:00Z&page=1&pageSize=10",
        headers=api_headers | {"X-Request-ID": "req-py-e2e-md-changes"},
    )
    assert changes.status_code == 200
    assert "records" in changes.json()["data"]

    health = client.get(
        "/api/v1/master-data/health",
        headers=api_headers | {"X-Request-ID": "req-py-e2e-md-health"},
    )
    assert health.status_code == 200
    assert health.json()["data"]["status"] == "ok"
    assert health.json()["data"]["database"] == "ok"

    with SessionLocal() as db:
        log_count = db.scalar(select(func.count()).select_from(ApiCallLog).where(ApiCallLog.trace_id.like("req-py-e2e-md-%")))
        client_row = db.get(ApiClient, "HOSPITAL_DATA_BUS")
        assert log_count >= 7
        assert client_row is not None
        assert client_row.last_used_at is not None


@requires_db
def test_nhsa_disabled_intersection_sets_material_inactive(tmp_path: Path) -> None:
    _clean_e2e_data()
    client = TestClient(app)
    mat_path = REPO_ROOT / "data/templates/H-UDMP-TEMPLATE-MATERIALS-v1.0.csv"

    mat_resp = _post_file(
        client,
        "/api/v1/materials/import",
        mat_path,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-DISABLE-SEED-001", "source_type": "MVP_TEMPLATE"},
    )
    assert mat_resp.status_code == 200
    assert mat_resp.json()["data"]["success_count"] == 2

    disabled_path = tmp_path / "disabled-intersection.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Query1"
    ws.append(["27位码", "原码状态"])
    ws.append(["123456789012345678901234567", "停用"])
    ws.append(["123456789012345678901234567", "停用"])
    wb.save(disabled_path)

    disabled_resp = _post_file(
        client,
        "/api/v1/materials/import",
        disabled_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-DISABLE-INTERSECTION-001",
            "source_type": "NHSA_DISABLED",
            "sheet_name": "Query1",
        },
    )
    assert disabled_resp.status_code == 200
    data = disabled_resp.json()["data"]
    assert data["source_row_count"] == 2
    assert data["unique_key_count"] == 1
    assert data["skipped_duplicate_count"] == 1
    assert data["success_count"] == 1
    assert data["disabled_count"] == 1

    material_search = client.get(
        "/api/v1/materials/search?yb_code_27=123456789012345678901234567&status=INACTIVE",
        headers=HEADERS,
    )
    assert material_search.status_code == 200
    assert material_search.json()["data"]["page"]["total"] == 1


@requires_db
def test_import_exception_scenarios(tmp_path: Path) -> None:
    _clean_e2e_data()
    client = TestClient(app)

    missing_dept_code = tmp_path / "departments-missing-code.csv"
    missing_dept_code.write_text("dept_name,dept_alias,status\n缺编码科室,,ACTIVE\n", encoding="utf-8")
    dept_resp = _post_file(
        client,
        "/api/v1/departments/import",
        missing_dept_code,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-IMPORT-ERR-DEPT-001"},
    )
    assert dept_resp.status_code == 200
    dept_data = dept_resp.json()["data"]
    assert dept_data["success_count"] == 0
    assert dept_data["failed_count"] == 1
    assert dept_data["failures"][0]["field"] == "dept_code"

    bad_source_type = tmp_path / "materials-invalid-source.csv"
    bad_source_type.write_text(
        "yb_code_27,generic_name,status\n123456789012345678901234567,测试耗材,ACTIVE\n",
        encoding="utf-8",
    )
    bad_source_resp = _post_file(
        client,
        "/api/v1/materials/import",
        bad_source_type,
        {"source_system": "PY-E2E", "source_tx_id": "PY-E2E-IMPORT-ERR-SOURCE-001", "source_type": "UNKNOWN"},
    )
    assert bad_source_resp.status_code == 400
    assert bad_source_resp.json()["code"] == "INVALID_SOURCE_TYPE"

    wrong_sheet_path = tmp_path / "materials-wrong-sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Query1"
    ws.append(["yb_code_27", "generic_name", "status"])
    ws.append(["123456789012345678901234567", "测试耗材", "ACTIVE"])
    wb.save(wrong_sheet_path)
    wrong_sheet_resp = _post_file(
        client,
        "/api/v1/materials/import",
        wrong_sheet_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-IMPORT-ERR-SHEET-001",
            "source_type": "MVP_TEMPLATE",
            "sheet_name": "MissingSheet",
        },
    )
    assert wrong_sheet_resp.status_code == 400
    assert wrong_sheet_resp.json()["code"] == "IMPORT_FILE_INVALID"
    assert "MissingSheet" in wrong_sheet_resp.json()["message"]

    duplicate_materials = tmp_path / "materials-duplicate.csv"
    duplicate_materials.write_text(
        "\n".join(
            [
                "yb_code_27,generic_name,spec_value,status",
                "123456789012345678901234567,测试耗材A,22mm,ACTIVE",
                "123456789012345678901234567,测试耗材B,24mm,ACTIVE",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    duplicate_resp = _post_file(
        client,
        "/api/v1/materials/import",
        duplicate_materials,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-IMPORT-ERR-DUP-001",
            "source_type": "MVP_TEMPLATE",
        },
    )
    assert duplicate_resp.status_code == 200
    duplicate_data = duplicate_resp.json()["data"]
    assert duplicate_data["source_row_count"] == 2
    assert duplicate_data["unique_key_count"] == 1
    assert duplicate_data["skipped_duplicate_count"] == 1
    assert duplicate_data["success_count"] == 1
    assert duplicate_data["duplicate_count"] == 1

    invalid_material_code = tmp_path / "materials-invalid-code.csv"
    invalid_material_code.write_text("yb_code_27,generic_name,status\n123,短码耗材,ACTIVE\n", encoding="utf-8")
    invalid_code_resp = _post_file(
        client,
        "/api/v1/materials/import",
        invalid_material_code,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-IMPORT-ERR-CODE-001",
            "source_type": "MVP_TEMPLATE",
        },
    )
    assert invalid_code_resp.status_code == 200
    invalid_code_data = invalid_code_resp.json()["data"]
    assert invalid_code_data["success_count"] == 0
    assert invalid_code_data["failed_count"] == 1
    assert invalid_code_data["failures"][0]["field"] == "yb_code_27"

    invalid_docx = tmp_path / "invalid-device-catalog.docx"
    invalid_docx.write_bytes(b"not a docx package")
    invalid_docx_resp = _post_file(
        client,
        "/api/v1/materials/import",
        invalid_docx,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-IMPORT-ERR-DOCX-001",
            "source_type": "DEVICE_CLASSIFICATION_CATALOG",
        },
    )
    assert invalid_docx_resp.status_code == 400
    assert invalid_docx_resp.json()["code"] == "IMPORT_FILE_INVALID"


@requires_db
def test_async_material_import_task_and_failures(tmp_path: Path) -> None:
    _clean_e2e_data()
    client = TestClient(app)
    mat_path = REPO_ROOT / "data/templates/H-UDMP-TEMPLATE-MATERIALS-v1.0.csv"

    accepted = _post_file(
        client,
        "/api/v1/import-tasks/materials",
        mat_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-ASYNC-MAT-001",
            "source_type": "MVP_TEMPLATE",
        },
    )
    assert accepted.status_code == 200
    assert accepted.json()["data"]["batch_id"] == "PY-E2E-ASYNC-MAT-001"

    status_resp = client.get("/api/v1/import-tasks/PY-E2E-ASYNC-MAT-001", headers=HEADERS)
    assert status_resp.status_code == 200
    status_data = status_resp.json()["data"]
    assert status_data["status"] == "COMPLETED"
    assert status_data["progress_percent"] == 100
    assert status_data["success_count"] == 2
    assert status_data["failed_count"] == 0
    assert status_data["source_file_name"] == mat_path.name
    assert status_data["source_file_extension"] == ".csv"
    assert status_data["source_file_size_bytes"] == mat_path.stat().st_size
    assert status_data["source_artifact"]["upload_mode"] == "single_file"
    SessionLocal = get_session_factory()
    with SessionLocal() as db:
        batch = db.get(SysImportBatch, "PY-E2E-ASYNC-MAT-001")
        assert batch is not None
        assert batch.source_file_path is not None
        assert not Path(batch.source_file_path).exists()

    catalog_path = tmp_path / "async-device-catalog.docx"
    _write_device_catalog_docx(catalog_path)
    catalog_accepted = _post_file(
        client,
        "/api/v1/import-tasks/materials",
        catalog_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-ASYNC-DEVICE-CATALOG-001",
            "source_type": "DEVICE_CLASSIFICATION_CATALOG",
        },
    )
    assert catalog_accepted.status_code == 200

    catalog_status = client.get("/api/v1/import-tasks/PY-E2E-ASYNC-DEVICE-CATALOG-001", headers=HEADERS)
    assert catalog_status.status_code == 200
    catalog_data = catalog_status.json()["data"]
    assert catalog_data["status"] == "COMPLETED"
    assert catalog_data["source_row_count"] == 3
    assert catalog_data["success_count"] == 3
    with SessionLocal() as db:
        assert (
            db.scalar(
                text(
                    "select count(*) from ref_device_classification_catalog "
                    "where batch_id='PY-E2E-ASYNC-DEVICE-CATALOG-001'"
                )
            )
            == 3
        )

    invalid_path = tmp_path / "async-invalid-materials.csv"
    invalid_path.write_text("yb_code_27,generic_name,status\n123,短码耗材,ACTIVE\n", encoding="utf-8")
    failed = _post_file(
        client,
        "/api/v1/import-tasks/materials",
        invalid_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-ASYNC-MAT-FAIL-001",
            "source_type": "MVP_TEMPLATE",
        },
    )
    assert failed.status_code == 200

    failed_status = client.get("/api/v1/import-tasks/PY-E2E-ASYNC-MAT-FAIL-001", headers=HEADERS)
    assert failed_status.status_code == 200
    failed_data = failed_status.json()["data"]
    assert failed_data["status"] == "COMPLETED_WITH_ERRORS"
    assert failed_data["failed_count"] == 1

    failures = client.get("/api/v1/import-tasks/PY-E2E-ASYNC-MAT-FAIL-001/failures", headers=HEADERS)
    assert failures.status_code == 200
    assert failures.json()["data"]["page"]["total"] == 1
    default_failure = failures.json()["data"]["items"][0]
    assert default_failure["field_name"] == "yb_code_27"
    assert default_failure["raw_payload_included"] is False
    assert "raw_payload" not in default_failure

    raw_failures = client.get(
        "/api/v1/import-tasks/PY-E2E-ASYNC-MAT-FAIL-001/failures?include_raw_payload=true",
        headers=HEADERS,
    )
    assert raw_failures.status_code == 200
    raw_failure = raw_failures.json()["data"]["items"][0]
    assert raw_failure["raw_payload_included"] is True
    assert raw_failure["raw_payload"]["field"] == "yb_code_27"

    steward_raw_failures = client.get(
        "/api/v1/import-tasks/PY-E2E-ASYNC-MAT-FAIL-001/failures?include_raw_payload=true",
        headers=HEADERS | {"X-Operator-Name": "data-steward", "X-Operator-Role": "data_steward"},
    )
    assert steward_raw_failures.status_code == 403
    assert steward_raw_failures.json()["code"] == "FORBIDDEN"


@requires_db
def test_manufacturer_vendor_master_data_flow(tmp_path: Path) -> None:
    _clean_e2e_data()
    client = TestClient(app)

    create_resp = client.post(
        "/api/v1/manufacturer-vendors",
        headers=HEADERS,
        json={
            "organization_code": "PY-E2E-VENDOR-001",
            "standard_name": "示例医疗器械有限公司",
            "english_name": "EXAMPLE MEDICAL DEVICE CO., LTD.",
            "short_name": "示例医疗",
            "alias_names": ["示例器械", "Example Medical"],
            "unified_social_credit_code": "91310000MA1K000001",
            "organization_type": "企业",
            "country_region": "中国",
            "data_source": "人工维护",
            "remark": "PY-E2E-vendor",
        },
    )
    assert create_resp.status_code == 200
    vendor = create_resp.json()["data"]
    vendor_id = vendor["id"]
    assert vendor["standard_name"] == "示例医疗器械有限公司"
    assert "示例器械" in vendor["alias_names"]

    duplicate_name = client.post(
        "/api/v1/manufacturer-vendors",
        headers=HEADERS,
        json={"organization_code": "PY-E2E-VENDOR-DUP-NAME", "standard_name": "示例医疗器械有限公司"},
    )
    assert duplicate_name.status_code == 409
    assert duplicate_name.json()["code"] == "DUPLICATE_STANDARD_NAME"

    duplicate_credit = client.post(
        "/api/v1/manufacturer-vendors",
        headers=HEADERS,
        json={
            "organization_code": "PY-E2E-VENDOR-DUP-CREDIT",
            "standard_name": "另一个医疗器械有限公司",
            "unified_social_credit_code": "91310000MA1K000001",
        },
    )
    assert duplicate_credit.status_code == 409
    assert duplicate_credit.json()["code"] == "DUPLICATE_CREDIT_CODE"

    for role_type, business_domain in [
        ("manufacturer", "consumable"),
        ("registrant", "device"),
        ("after_sales", "equipment"),
        ("supplier", "service"),
    ]:
        role_resp = client.post(
            f"/api/v1/manufacturer-vendors/{vendor_id}/roles",
            headers=HEADERS,
            json={
                "role_type": role_type,
                "business_domain": business_domain,
                "remark": "PY-E2E-role",
            },
        )
        assert role_resp.status_code == 200

    child_resp = client.post(
        "/api/v1/manufacturer-vendors",
        headers=HEADERS,
        json={
            "organization_code": "PY-E2E-VENDOR-CHILD",
            "standard_name": "示例医疗器械子公司",
            "short_name": "示例子公司",
            "remark": "PY-E2E-vendor",
        },
    )
    assert child_resp.status_code == 200
    child_id = child_resp.json()["data"]["id"]

    relation_resp = client.post(
        "/api/v1/manufacturer-vendor-relations",
        headers=HEADERS,
        json={
            "parent_org_id": vendor_id,
            "child_org_id": child_id,
            "relation_type": "subsidiary",
            "effective_date": "2026-01-01",
            "remark": "PY-E2E-relation",
        },
    )
    assert relation_resp.status_code == 200
    assert relation_resp.json()["data"]["relation_name"] == "子公司"

    rename_resp = client.post(
        "/api/v1/manufacturer-vendor-relations",
        headers=HEADERS,
        json={
            "parent_org_id": vendor_id,
            "child_org_id": child_id,
            "relation_type": "renamed_from",
            "remark": "PY-E2E-relation",
        },
    )
    assert rename_resp.status_code == 200

    mapping_resp = client.post(
        f"/api/v1/manufacturer-vendors/{vendor_id}/external-mappings",
        headers=HEADERS,
        json={
            "system_name": "NHSA",
            "external_code": "PY-E2E-NHSA-VENDOR-001",
            "external_name": "示例器械",
            "confidence": 1,
        },
    )
    assert mapping_resp.status_code == 200

    alias_search = client.get("/api/external/manufacturer-vendors?keyword=示例器械", headers=HEADERS)
    assert alias_search.status_code == 200
    assert alias_search.json()["data"]["page"]["total"] >= 1

    credit_search = client.get("/api/external/manufacturer-vendors?keyword=91310000MA1K000001", headers=HEADERS)
    assert credit_search.status_code == 200
    assert credit_search.json()["data"]["items"][0]["id"] == vendor_id

    detail_resp = client.get(f"/api/external/manufacturer-vendors/{vendor_id}", headers=HEADERS)
    assert detail_resp.status_code == 200
    detail = detail_resp.json()["data"]
    assert len(detail["roles"]) == 4
    assert len(detail["external_mappings"]) == 1
    assert len(detail["relations"]) >= 2

    relation_list = client.get(f"/api/external/manufacturer-vendors/{vendor_id}/relations", headers=HEADERS)
    assert relation_list.status_code == 200
    assert relation_list.json()["data"]["subsidiaries"]

    nhsa_file = tmp_path / "nhsa-vendor.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Query1"
    ws.append(["27位码", "20位码", "医保通用名", "单件产品名称", "耗材企业", "注册备案号", "规格", "型号"])
    ws.append(["323456789012345678901234567", "32345678901234567890", "测试导管", "测试导管", "示例器械", "国械注准2026TEST", "22mm", "A"])
    ws.append(["323456789012345678901234568", "32345678901234567890", "候选导管", "候选导管", "候选厂商有限公司", "国械注准2026TEST2", "24mm", "B"])
    wb.save(nhsa_file)

    import_resp = _post_file(
        client,
        "/api/v1/materials/import",
        nhsa_file,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-NHSA-VENDOR-IMPORT",
            "source_type": "NHSA_FULL_SPEC",
            "sheet_name": "Query1",
        },
    )
    assert import_resp.status_code == 200
    assert import_resp.json()["data"]["vendor_candidate_count"] == 2

    with get_session_factory()() as db:
        matched_material = db.scalar(
            text(
                "select manufacturer_org_id from dict_material_specs "
                "where yb_code_27='323456789012345678901234567'"
            )
        )
        assert str(matched_material) == vendor_id
        candidate_id = db.scalar(
            text(
                "select id from manufacturer_vendor_candidate "
                "where normalized_name='候选厂商有限公司' and source_system='NHSA'"
            )
        )
        assert candidate_id is not None

    candidate_list = client.get("/api/v1/manufacturer-vendor-candidates?source_system=NHSA", headers=HEADERS)
    assert candidate_list.status_code == 200
    assert candidate_list.json()["data"]["page"]["total"] >= 2

    map_candidate = client.post(
        f"/api/v1/manufacturer-vendor-candidates/{candidate_id}/map",
        headers=HEADERS,
        json={"org_id": vendor_id, "reviewer": "data_admin"},
    )
    assert map_candidate.status_code == 200
    assert map_candidate.json()["data"]["match_status"] == "matched"

    merge_candidate_path = tmp_path / "merge-candidate.csv"
    merge_candidate_path.write_text(
        "yb_code_27,generic_name,manufacturer,status\n423456789012345678901234568,合并候选耗材,示例器械旧称,ACTIVE\n",
        encoding="utf-8",
    )
    merge_import = _post_file(
        client,
        "/api/v1/materials/import",
        merge_candidate_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-MERGE-VENDOR-CANDIDATE",
            "source_type": "MVP_TEMPLATE",
        },
    )
    assert merge_import.status_code == 200

    with get_session_factory()() as db:
        merge_candidate_id = db.scalar(
            text("select id from manufacturer_vendor_candidate where normalized_name='示例器械旧称'")
        )
        assert merge_candidate_id is not None

    merge_candidate = client.post(
        f"/api/v1/manufacturer-vendor-candidates/{merge_candidate_id}/merge",
        headers=HEADERS,
        json={"org_id": vendor_id, "reviewer": "data_admin"},
    )
    assert merge_candidate.status_code == 200
    assert merge_candidate.json()["data"]["suggested_action"] == "merge"

    alias_after_merge = client.get("/api/external/manufacturer-vendors?keyword=示例器械旧称", headers=HEADERS)
    assert alias_after_merge.status_code == 200
    assert alias_after_merge.json()["data"]["page"]["total"] >= 1

    manual_candidate_path = tmp_path / "manual-candidate.csv"
    manual_candidate_path.write_text(
        "yb_code_27,generic_name,manufacturer,status\n423456789012345678901234567,新候选耗材,新候选厂商有限公司,ACTIVE\n",
        encoding="utf-8",
    )
    manual_import = _post_file(
        client,
        "/api/v1/materials/import",
        manual_candidate_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-MANUAL-VENDOR-CANDIDATE",
            "source_type": "MVP_TEMPLATE",
        },
    )
    assert manual_import.status_code == 200

    with get_session_factory()() as db:
        new_candidate_id = db.scalar(
            text("select id from manufacturer_vendor_candidate where normalized_name='新候选厂商有限公司'")
        )
        assert new_candidate_id is not None

    create_from_candidate = client.post(
        f"/api/v1/manufacturer-vendor-candidates/{new_candidate_id}/create-vendor",
        headers=HEADERS,
        json={
            "reviewer": "data_admin",
            "vendor": {"organization_code": "PY-E2E-VENDOR-CANDIDATE", "standard_name": "新候选厂商有限公司"},
            "roles": [{"role_type": "manufacturer", "business_domain": "consumable"}],
        },
    )
    assert create_from_candidate.status_code == 200
    assert create_from_candidate.json()["data"]["candidate"]["match_status"] == "matched"

    batch_candidate_path = tmp_path / "batch-candidate.csv"
    batch_candidate_path.write_text(
        "yb_code_27,generic_name,manufacturer,status\n423456789012345678901234569,批量候选耗材,批量忽略厂商,ACTIVE\n",
        encoding="utf-8",
    )
    batch_import = _post_file(
        client,
        "/api/v1/materials/import",
        batch_candidate_path,
        {
            "source_system": "PY-E2E",
            "source_tx_id": "PY-E2E-BATCH-VENDOR-CANDIDATE",
            "source_type": "MVP_TEMPLATE",
        },
    )
    assert batch_import.status_code == 200

    with get_session_factory()() as db:
        batch_candidate_id = db.scalar(
            text("select id from manufacturer_vendor_candidate where normalized_name='批量忽略厂商'")
        )
        assert batch_candidate_id is not None

    batch_review = client.post(
        "/api/v1/manufacturer-vendor-candidates/batch",
        headers=HEADERS,
        json={"actions": [{"candidate_id": str(batch_candidate_id), "action": "ignore", "reviewer": "data_admin"}]},
    )
    assert batch_review.status_code == 200
    assert batch_review.json()["data"]["items"][0]["status"] == "ok"

    equipment_reference = client.get(f"/api/external/manufacturer-vendors/{vendor_id}", headers=HEADERS)
    assert equipment_reference.status_code == 200
    assert equipment_reference.json()["data"]["id"] == vendor_id


@requires_db
def test_equipment_dictionary_real_database_contract() -> None:
    client = TestClient(app)

    with get_session_factory()() as db:
        category_count = db.scalar(text("select count(*) from dict_equipment_categories where status='ACTIVE'"))
        standard_name_count = db.scalar(text("select count(*) from dict_equipment_standard_names where status='ACTIVE'"))
        device_classification_count = db.scalar(
            text("select count(*) from ref_device_classification_catalog where data_status='effective' and hidden_in_tree=false")
        )

    checks = [
        ("/api/v1/equipment/categories?page=1&page_size=1", category_count),
        ("/api/v1/equipment/standard-names?page=1&page_size=1", standard_name_count),
        ("/api/v1/equipment/device-classifications?page=1&page_size=1", device_classification_count),
        ("/api/external/equipment/categories?page=1&page_size=1", category_count),
        ("/api/external/equipment/standard-names?page=1&page_size=1", standard_name_count),
        ("/api/external/equipment/device-classifications?page=1&page_size=1", device_classification_count),
    ]
    for path, expected_total in checks:
        response = client.get(path, headers=HEADERS)
        assert response.status_code == 200
        body = response.json()
        assert body["success"] is True
        assert body["data"]["page"]["total"] == expected_total
        assert len(body["data"]["items"]) <= 1
