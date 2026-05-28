import zipfile

import py7zr
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core.config import get_settings
from app.main import app


def test_health_endpoint() -> None:
    client = TestClient(app)
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_frontend_origin_can_call_health() -> None:
    client = TestClient(app)
    response = client.options(
        "/health",
        headers={
            "Origin": "http://127.0.0.1:5101",
            "Access-Control-Request-Method": "GET",
        },
    )
    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://127.0.0.1:5101"


def test_frontend_origin_can_send_session_token_header() -> None:
    client = TestClient(app)
    response = client.options(
        "/api/v1/auth/me",
        headers={
            "Origin": "http://127.0.0.1:5101",
            "Access-Control-Request-Method": "GET",
            "Access-Control-Request-Headers": "X-Session-Token,Content-Type",
        },
    )
    assert response.status_code == 200
    allow_headers = response.headers["access-control-allow-headers"]
    assert "X-Session-Token" in allow_headers


def test_frontend_origin_can_patch_department_status() -> None:
    client = TestClient(app)
    response = client.options(
        "/api/v1/departments/NHC-03.04/status",
        headers={
            "Origin": "http://127.0.0.1:5101",
            "Access-Control-Request-Method": "PATCH",
            "Access-Control-Request-Headers": "X-Session-Token,Content-Type",
        },
    )
    assert response.status_code == 200
    assert "PATCH" in response.headers["access-control-allow-methods"]


def test_openapi_contains_mvp_paths() -> None:
    client = TestClient(app)
    paths = client.get("/openapi.json").json()["paths"]
    expected = {
        "/api/v1/departments/search",
        "/api/v1/auth/login",
        "/api/v1/auth/me",
        "/api/v1/auth/permissions",
        "/api/v1/auth/operators",
        "/api/v1/dictionaries/catalog",
        "/api/v1/departments/import",
        "/api/v1/departments/import/inspect",
        "/api/v1/departments/standard-library/sync",
        "/api/v1/departments/{dept_code}/status",
        "/api/v1/materials/search",
        "/api/v1/materials/import",
        "/api/v1/materials/import/inspect",
        "/api/v1/materials/transcode/resolve",
        "/api/v1/import-tasks",
        "/api/v1/import-tasks/materials",
        "/api/v1/import-tasks/materials/archive",
        "/api/v1/import-tasks/materials/archive/inspect",
        "/api/v1/import-tasks/{batch_id}",
        "/api/v1/import-tasks/{batch_id}/failures",
        "/api/v1/equipment/categories",
        "/api/v1/equipment/standard-names",
        "/api/v1/equipment/device-classifications",
        "/api/v1/equipment/device-classifications/export/templates",
        "/api/v1/equipment/device-classifications/export/estimate",
        "/api/v1/equipment/device-classifications/export/tasks",
        "/api/v1/equipment/device-classifications/export/tasks/{task_no}",
        "/api/v1/equipment/device-classifications/export/tasks/{task_no}/download",
        "/api/v1/equipment/device-classifications/corrections",
        "/api/v1/equipment/device-classifications/{catalog_id}/corrections",
        "/api/v1/equipment/device-classifications/corrections/{correction_id}/approve",
        "/api/v1/equipment/device-classifications/corrections/{correction_id}/reject",
        "/api/v1/equipment/device-classifications/corrections/{correction_id}/execute",
        "/api/v1/equipment/device-classifications/{catalog_id}/governance",
        "/api/v1/equipment/device-classifications/{catalog_id}/impact-analysis",
        "/api/v1/equipment/device-classifications/{catalog_id}/source-file/download",
        "/api/v1/equipment/import",
        "/api/v1/master-data/device-classification/catalog",
        "/api/v1/master-data/device-classification/catalog/{id}",
        "/api/v1/master-data/device-classification/catalog/by-code/{code}",
        "/api/v1/master-data/device-classification/tree",
        "/api/v1/master-data/device-classification/current-version",
        "/api/v1/master-data/device-classification/changes",
        "/api/v1/master-data/health",
        "/api/external/equipment/categories",
        "/api/external/equipment/standard-names",
        "/api/external/equipment/device-classifications",
        "/api/v1/mapping/resolve",
        "/api/v1/mapping/review-tasks",
        "/api/v1/mapping/review-tasks/{task_id}/approve",
        "/api/v1/mapping/review-tasks/{task_id}/reject",
        "/api/v1/exchange/logs",
        "/api/v1/exchange/logs/{log_id}",
        "/health",
    }
    assert expected.issubset(set(paths))


def test_dictionary_catalog_publishes_available_and_blueprint_dictionaries() -> None:
    client = TestClient(app)
    response = client.get("/api/v1/dictionaries/catalog", headers={"X-API-Key": "change-me"})

    assert response.status_code == 200
    data = response.json()["data"]
    items = {item["dictionary_key"]: item for item in data["items"]}
    assert data["summary"]["available_count"] >= 2
    assert items["DEPARTMENT"]["status"] == "AVAILABLE"
    assert items["DEPARTMENT"]["key_fields"][0]["name"] == "dept_code"
    assert items["DRUG"]["status"] == "BLUEPRINT"
    assert {field["name"] for field in items["CAMPUS"]["key_fields"]} >= {"campus_code", "campus_name"}
    assert items["CAMPUS"]["status"] == "AVAILABLE"
    assert items["UNIT"]["status"] == "AVAILABLE"
    assert len(items["UNIT"]["reference_records"]) >= 10
    assert items["SUPPLIER"]["candidate_source"].startswith("医保耗材")
    assert items["WARD"]["reference_records"][0]["ward_name"]
    assert items["DIAGNOSIS"]["status"] == "AVAILABLE"
    assert items["PROCEDURE"]["status"] == "AVAILABLE"


def test_protected_api_uses_error_envelope_without_api_key() -> None:
    client = TestClient(app)
    response = client.get("/api/v1/materials/search")
    assert response.status_code == 401
    body = response.json()
    assert body["success"] is False
    assert body["code"] == "UNAUTHORIZED"
    assert body["data"] is None


def test_operator_role_permission_is_enforced() -> None:
    client = TestClient(app)
    response = client.get(
        "/api/v1/materials/search",
        headers={"X-API-Key": "change-me", "X-Operator-Name": "uat-auditor", "X-Operator-Role": "auditor"},
    )
    assert response.status_code == 403
    body = response.json()
    assert body["success"] is False
    assert body["code"] == "FORBIDDEN"


def test_auth_me_returns_operator_permissions() -> None:
    client = TestClient(app)
    response = client.get(
        "/api/v1/auth/me",
        headers={
            "X-API-Key": "change-me",
            "X-Operator-Name": "data-admin",
            "X-Operator-Role": "data_steward",
        },
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["operator_name"] == "data-admin"
    assert data["operator_role"] == "data_steward"
    assert "mapping.review" in data["permissions"]
    assert "exchange.view" not in data["permissions"]
    assert "raw_payload.view" not in data["permissions"]


def test_material_import_inspect_lists_excel_sheets(tmp_path) -> None:
    path = tmp_path / "materials.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "封面"
    ws.append(["note"])
    ws.append(["ignore"])
    ws2 = wb.create_sheet("全量规格型号信息")
    ws2.append(["yb_code_27", "generic_name", "status"])
    ws2.append(["123456789012345678901234567", "测试耗材", "ACTIVE"])
    wb.save(path)

    client = TestClient(app)
    with path.open("rb") as handle:
        response = client.post(
            "/api/v1/materials/import/inspect",
            headers={"X-API-Key": "change-me", "X-Operator-Name": "data-admin", "X-Operator-Role": "data_steward"},
            data={"source_type": "MVP_TEMPLATE"},
            files={"file": ("materials.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["file_type"] == "EXCEL"
    assert data["sheet_required"] is True
    assert [sheet["name"] for sheet in data["sheets"]] == ["封面", "全量规格型号信息"]
    assert data["default_sheet_name"] == "封面"


def test_material_archive_inspect_identifies_importable_files(tmp_path) -> None:
    archive_path = tmp_path / "nhsa.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr(
            "12、医保医用耗材代码_C05_心脏外科类材料_全量规格型号信息.xlsx",
            _xlsx_bytes("Query1", ["27位码", "20位码", "医保通用名"], ["123456789012345678901234567", "12345678901234567890", "测试耗材"]),
        )
        archive.writestr(
            "30、停用表.xlsx",
            _xlsx_bytes("Query1", ["27位码", "原码状态"], ["123456789012345678901234567", "停用"]),
        )
        archive.writestr("说明.txt", "readme")

    client = TestClient(app)
    with archive_path.open("rb") as handle:
        response = client.post(
            "/api/v1/import-tasks/materials/archive/inspect",
            headers={"X-API-Key": "change-me", "X-Operator-Name": "data-admin", "X-Operator-Role": "data_steward"},
            files={"file": ("nhsa.zip", handle, "application/zip")},
        )

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["archive_id"].startswith("ARCHIVE-")
    assert data["archive_file_name"] == "nhsa.zip"
    assert data["archive_file_extension"] == ".zip"
    assert data["archive_file_size_bytes"] == archive_path.stat().st_size
    assert data["total_uncompressed_size_bytes"] > 0
    assert data["importable_count"] == 2
    assert data["source_type_counts"] == {"NHSA_FULL_SPEC": 1, "NHSA_DISABLED": 1}
    by_name = {item["file_name"]: item for item in data["items"]}
    assert by_name["12、医保医用耗材代码_C05_心脏外科类材料_全量规格型号信息.xlsx"]["source_type"] == "NHSA_FULL_SPEC"
    assert by_name["12、医保医用耗材代码_C05_心脏外科类材料_全量规格型号信息.xlsx"]["extension"] == ".xlsx"
    assert by_name["12、医保医用耗材代码_C05_心脏外科类材料_全量规格型号信息.xlsx"]["row_count"] == 1
    assert by_name["30、停用表.xlsx"]["source_type"] == "NHSA_DISABLED"
    assert by_name["说明.txt"]["supported"] is False


def test_material_archive_inspect_supports_7z(tmp_path) -> None:
    payload_path = tmp_path / "31、转码表.xlsx"
    payload_path.write_bytes(
        _xlsx_bytes("Query1", ["原27位码", "原码状态", "27位码", "类型"], ["123456789012345678901234567", "停用", "223456789012345678901234567", "转码"])
    )
    archive_path = tmp_path / "nhsa.7z"
    with py7zr.SevenZipFile(archive_path, "w") as archive:
        archive.write(payload_path, arcname=payload_path.name)

    client = TestClient(app)
    with archive_path.open("rb") as handle:
        response = client.post(
            "/api/v1/import-tasks/materials/archive/inspect",
            headers={"X-API-Key": "change-me", "X-Operator-Name": "data-admin", "X-Operator-Role": "data_steward"},
            files={"file": ("nhsa.7z", handle, "application/x-7z-compressed")},
        )

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["archive_format"] == "7Z"
    assert data["archive_file_extension"] == ".7z"
    assert data["importable_count"] == 1
    assert data["items"][0]["source_type"] == "NHSA_TRANSCODE"


def _xlsx_bytes(sheet_name: str, headers: list[str], row: list[str]) -> bytes:
    from io import BytesIO

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(headers)
    worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_operator_password_login_returns_session_token() -> None:
    client = TestClient(app)
    login = client.post("/api/v1/auth/login", json={"username": "E1001", "password": "demo123"})
    assert login.status_code == 200
    data = login.json()["data"]
    assert data["operator_name"] == "E1001"
    assert data["operator_role"] == "data_steward"
    assert data["session_token"]
    assert data["session_expires_at"]
    assert "mapping.review" in data["permissions"]

    me = client.get("/api/v1/auth/me", headers={"X-Session-Token": data["session_token"]})
    assert me.status_code == 200
    me_data = me.json()["data"]
    assert me_data["operator_name"] == "E1001"
    assert me_data["operator_role"] == "data_steward"
    assert me_data["auth_model"] == "operator_password_session"
    assert me_data["session_expires_at"] == data["session_expires_at"]


def test_expired_operator_session_returns_session_error(monkeypatch) -> None:
    monkeypatch.setenv("OPERATOR_SESSION_TTL_MINUTES", "-1")
    get_settings.cache_clear()
    client = TestClient(app)
    login = client.post("/api/v1/auth/login", json={"username": "E1001", "password": "demo123"})
    token = login.json()["data"]["session_token"]

    response = client.get("/api/v1/auth/me", headers={"X-Session-Token": token})
    get_settings.cache_clear()

    assert response.status_code == 401
    body = response.json()
    assert body["success"] is False
    assert body["code"] == "SESSION_EXPIRED"
    assert body["message"] == "登录已过期，请重新登录"


def test_operator_password_login_rejects_bad_password() -> None:
    client = TestClient(app)
    response = client.post("/api/v1/auth/login", json={"username": "E1001", "password": "wrong"})
    assert response.status_code == 401
    assert response.json()["code"] == "UNAUTHORIZED"


def test_import_upload_too_large_returns_error_envelope(monkeypatch) -> None:
    monkeypatch.setenv("HUDMP_IMPORT_MAX_UPLOAD_BYTES", "8")
    get_settings.cache_clear()
    client = TestClient(app)
    response = client.post(
        "/api/v1/departments/import",
        headers={"X-API-Key": "change-me"},
        data={"source_system": "PY-CONTRACT", "source_tx_id": "PY-CONTRACT-UPLOAD-LIMIT"},
        files={"file": ("too-large.csv", b"123456789", "text/csv")},
    )
    get_settings.cache_clear()

    assert response.status_code == 413
    body = response.json()
    assert body["success"] is False
    assert body["code"] == "UPLOAD_TOO_LARGE"
    assert body["data"] is None


def test_auth_me_platform_admin_can_view_raw_payload_permission() -> None:
    client = TestClient(app)
    response = client.get(
        "/api/v1/auth/me",
        headers={
            "X-API-Key": "change-me",
            "X-Operator-Name": "platform-admin",
            "X-Operator-Role": "platform_admin",
        },
    )
    assert response.status_code == 200
    assert "raw_payload.view" in response.json()["data"]["permissions"]


def test_auth_permissions_returns_backend_matrix() -> None:
    client = TestClient(app)
    response = client.get(
        "/api/v1/auth/permissions",
        headers={"X-API-Key": "change-me", "X-Operator-Name": "platform-admin", "X-Operator-Role": "platform_admin"},
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["source"] == "backend"
    roles = {row["role"]: set(row["permissions"]) for row in data["roles"]}
    assert "raw_payload.view" in roles["platform_admin"]
    assert "raw_payload.view" not in roles["data_steward"]
    assert "exchange.view" in roles["auditor"]


def test_auth_permissions_requires_permission_manager() -> None:
    client = TestClient(app)
    login = client.post("/api/v1/auth/login", json={"username": "E1001", "password": "demo123"})
    token = login.json()["data"]["session_token"]

    response = client.get("/api/v1/auth/permissions", headers={"X-Session-Token": token})

    assert response.status_code == 403
    assert response.json()["code"] == "FORBIDDEN"


def test_auth_operators_returns_preconfigured_accounts_without_passwords() -> None:
    client = TestClient(app)
    response = client.get(
        "/api/v1/auth/operators",
        headers={"X-API-Key": "change-me", "X-Operator-Name": "platform-admin", "X-Operator-Role": "platform_admin"},
    )

    assert response.status_code == 200
    data = response.json()["data"]
    accounts = {account["username"]: account for account in data["accounts"]}
    assert data["source"] == "backend"
    assert accounts["admin"]["role"] == "platform_admin"
    assert "permissions.manage" in accounts["admin"]["permissions"]
    assert accounts["E1001"]["display_name"] == "数据治理员"
    assert "password" not in accounts["admin"]


def test_unknown_operator_role_is_rejected() -> None:
    client = TestClient(app)
    response = client.get(
        "/api/v1/exchange/logs",
        headers={"X-API-Key": "change-me", "X-Operator-Name": "unknown", "X-Operator-Role": "superuser"},
    )
    assert response.status_code == 403
    assert response.json()["code"] == "FORBIDDEN"


def test_validation_error_uses_error_envelope() -> None:
    client = TestClient(app)
    response = client.get("/api/v1/materials/transcode/resolve", headers={"X-API-Key": "change-me"})
    assert response.status_code == 400
    body = response.json()
    assert body["success"] is False
    assert body["code"] == "VALIDATION_FAILED"
    assert body["data"] is None
    assert body["details"]
